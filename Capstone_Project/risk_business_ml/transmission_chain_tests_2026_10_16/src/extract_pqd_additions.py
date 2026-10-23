from pathlib import Path
import openpyxl,pandas as pd,json,hashlib
ROOT=Path(__file__).resolve().parents[1];OLD=ROOT.parent/'event_vulnerability_tests_2026_10_06'
rows=[];detail=[]
for f in sorted((ROOT/'inputs/pqd').glob('*.xlsx')):
 w=openpyxl.load_workbook(f,data_only=True,read_only=True)
 s=w['DTCC_ConsolidatedDataFile'];q=f.stem.replace('DTCC_PQD_','').replace('_','')
 for rn,row in enumerate(s.iter_rows(min_row=1,max_row=min(s.max_row,500),max_col=45,values_only=True),1):
  if len(row)<2 or str(row[1]) not in ['17.2.1','17.3.1','19.1.3.1','19.1.3.2']:continue
  code=str(row[1])
  # Headers changed in 2023. Resolve each group from its own ValueN header.
  if 'headers' not in locals() or last_source!=str(f):
   headers=list(next(s.iter_rows(min_row=1,max_row=1,max_col=45,values_only=True)));last_source=str(f)
  for label,group in [('GSD',1),('MBSD',2),('NSCC',3),('ALL',5)]:
   k=headers.index(f'Value{group}')-4
   raw=row[k+4];value=None
   try:value=float(str(raw).replace('%','').replace(',',''))
   except (ValueError,TypeError):pass
   rows.append({'quarter':q,'entity':label,'code':code,'value_as_reported':value,'raw_value':raw,'definition':row[3],'notes':row[k+5], 'reported_group_identifier':row[k], 'number_format':s.cell(rn,k+5).number_format,'source_file':f.name,'sheet':s.title,'cell':f'{openpyxl.utils.get_column_letter(k+5)}{rn}','scope':'source has shared-system caveat; do not assume independent entity measurement' if code.startswith('17.') else 'client-clearing concentration; peak and average separate','sha256':hashlib.sha256(f.read_bytes()).hexdigest()})
 if 'DTCC_DataFile_17_3' in w.sheetnames:
  s=w['DTCC_DataFile_17_3']
  for rn,row in enumerate(s.iter_rows(max_row=min(s.max_row,100),max_col=12,values_only=True),1):
   if any(x is not None for x in row):detail.append({'quarter':q,'source_file':f.name,'sheet':s.title,'row':rn,'values':json.dumps(row,default=str,ensure_ascii=False)})
 w.close()
a=pd.DataFrame(rows);a.to_csv(ROOT/'data/pqd_additional_cell_evidence.csv',index=False)
pd.DataFrame(detail).to_csv(ROOT/'data/pqd_17_3_detail_evidence.csv',index=False)
print(a.groupby(['code','entity']).agg(rows=('quarter','size'),numeric=('value_as_reported','count')).to_string())
print(a[(a.quarter=='2026Q2')][['entity','code','raw_value']].to_string(index=False))
print('Missing 2024Q4: original source is PDF; retained as a gap until verified extraction.')
