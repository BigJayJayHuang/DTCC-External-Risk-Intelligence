from pathlib import Path
from urllib.parse import urljoin
import json,re
import numpy as np
import pandas as pd
import openpyxl
from bs4 import BeautifulSoup

ROOT=Path(__file__).resolve().parents[1]
OLD=ROOT.parent
RAW=ROOT/'raw';DATA=ROOT/'data'
QUARTERS=pd.period_range('2019Q1','2026Q2',freq='Q')
families=[('business_continuity','业务连续性'),('cybersecurity','网络安全与信息保护'),('information_technology','信息技术'),
          ('data_management','数据管理'),('processing_operations','处理与运营'),('clearing_financial','清算与结算金融风险'),
          ('corporate_financial','公司财务'),('model_risk','模型风险'),('legal_regulatory','法律与监管合规'),
          ('third_party','第三方风险'),('human_capital','人员与人力资本'),('strategic','战略风险')]
pd.DataFrame(families,columns=['family','family_zh']).to_csv(DATA/'risk_taxonomy.csv',index=False)
records=[];definitions=[]

def add(feature,family,label,series,source,units,scope='external',entity='ALL',limit='',primary=False,window='quarter'):
    definitions.append({'feature':feature,'family':family,'label_zh':label,'source':source,'units':units,
                        'scope':scope,'entity_scope':entity,'interpretation_limit':limit,'primary':primary,'measurement_window':window})
    for q,v in series.items():
        if pd.notna(v): records.append({'quarter':str(q),'feature':feature,'entity_scope':entity,'value':float(v)})

# Existing validated financial proxies: preserve the original definitions.
old_path=DATA/'legacy_risk_quarterly.csv'
if not old_path.exists():old_path.write_bytes((OLD/'data/risk_quarterly.csv').read_bytes())
panel_path=DATA/'legacy_analysis_panel.csv'
if not panel_path.exists():panel_path.write_bytes((OLD/'data/analysis_panel.csv').read_bytes())
old=pd.read_csv(old_path).set_index('quarter')
old.index=pd.PeriodIndex(old.index,freq='Q')
for feature,label,units,source in [
 ('vix_change_pct','VIX季度均值增幅','percent','FRED VIXCLS'),
 ('rate_change_pp','十年期国债收益率季度末变化','percentage points','FRED DGS10'),
 ('credit_change_index','信用条件季度变化','index points','FRED NFCICREDIT'),
 ('funding_basis_max_bps','SOFR−EFFR季度最大利差','basis points','FRED SOFR/EFFR')]:
    add(feature,'clearing_financial',label,old.loc['2020Q1':'2026Q2',feature],source,units,
        limit='外部市场代理，不是会员违约、保证金不足或自身流动性缺口。',primary=feature=='funding_basis_max_bps')
for feature,family,label in [('fed_regulatory_announcements','legal_regulatory','美联储监管公告数量'),
                            ('fed_monetary_announcements','strategic','美联储货币政策公告数量')]:
    add(feature,family,label,old.loc['2020Q1':'2026Q2',feature],'Federal Reserve official press archives','count',limit='公告数量不是风险严重度；并非所有公告与DTCC有关。')

# FEMA county rows repeat declarations. Count a declaration once nationally.
rows=[]
for path in sorted(RAW.glob('fema_declarations_*.json')):
    rows.extend(json.loads(path.read_text()).get('DisasterDeclarationsSummaries',[]))
fema=pd.DataFrame(rows).drop_duplicates(['disasterNumber','declarationDate','state','declarationType'])
fema['date']=pd.to_datetime(fema.declarationDate,utc=True)
fema['quarter']=fema.date.dt.tz_localize(None).dt.to_period('Q')
fema=fema[fema.date.between(pd.Timestamp('2019-01-01',tz='UTC'),pd.Timestamp('2026-06-30 23:59:59',tz='UTC'))]
fema.to_csv(DATA/'fema_unique_declarations.csv',index=False)
major=fema[fema.declarationType=='DR']
add('fema_major_declarations','business_continuity','美国重大灾害宣告数量',major.groupby('quarter').disasterNumber.nunique().reindex(QUARTERS,fill_value=0),
    'OpenFEMA DisasterDeclarationsSummaries v2','declarations',limit='按宣告日期计数；全国外部中断触发代理，不证明DTCC服务中断。',primary=True)
add('fema_nonbiological_declarations','business_continuity','排除生物类的重大灾害宣告',major[major.incidentType!='Biological'].groupby('quarter').disasterNumber.nunique().reindex(QUARTERS,fill_value=0),
    'OpenFEMA DisasterDeclarationsSummaries v2','declarations',limit='用于检验疫情宣告是否主导结果，不衡量受灾损失。')

# KEV is a catalog of exploited vulnerabilities, not a list of company attacks.
kev=pd.DataFrame(json.loads((RAW/'cisa_kev.json').read_text())['vulnerabilities'])
assert not kev.cveID.duplicated().any()
kev['quarter']=pd.to_datetime(kev.dateAdded).dt.to_period('Q')
kev.to_csv(DATA/'cisa_catalog_entries.csv',index=False)
cyber_period=pd.period_range('2022Q1','2026Q2',freq='Q')
add('kev_new_entries','cybersecurity','CISA已利用漏洞目录新增数量',kev.groupby('quarter').size().reindex(cyber_period,fill_value=0),
    'CISA Known Exploited Vulnerabilities catalog','CVE entries',limit='按dateAdded计数；排除2021Q4目录启动批量载入。没有DTCC软件清单，不能证明漏洞暴露或攻击。',primary=True)

# Official product histories contain years absent from the recent JSON feed.
products=json.loads((ROOT/'history_products.json').read_text());cloud_rows=[]
for item in products:
    path=ROOT/item['file']
    if not path.exists():continue
    soup=BeautifulSoup(path.read_text(),'html.parser')
    for tr in soup.select('tr'):
        a=next((a for a in tr.select('a[href]') if 'incidents/' in a['href']),None)
        cells=tr.select('td')
        if not a or len(cells)<2:continue
        date_text=cells[0].get_text(' ',strip=True)
        date=pd.to_datetime(date_text,errors='coerce')
        if pd.isna(date):continue
        duration_text=cells[1].get_text(' ',strip=True)
        duration=sum(int(n)*{'day':24,'hour':1,'minute':1/60,'second':1/3600}[u.rstrip('s')]
                     for n,u in re.findall(r'(\d+)\s+(days?|hours?|minutes?|seconds?)',duration_text))
        cloud_rows.append({'group':item['group'],'product':item['product'],'date':date.strftime('%Y-%m-%d'),
                           'quarter':str(date.to_period('Q')),'incident_id':a['href'].split('/')[-1],
                           'incident_url':urljoin(item['url'],a['href']),'summary':(a.select_one('[class$="__summary-text"]') or a).get_text(' ',strip=True),
                           'duration_hours':duration,'duration_text':duration_text,'history_source':item['url']})
cloud=pd.DataFrame(cloud_rows)
cloud.to_csv(DATA/'provider_product_incidents.csv',index=False)
complete_cloud=pd.period_range('2022Q1','2026Q2',freq='Q')
for group,family,feature,label in [('it','information_technology','cloud_infrastructure_incidents','外部云计算与网络服务事件公告'),
                                  ('data','data_management','cloud_data_service_incidents','外部云数据服务事件公告'),
                                  ('third_party','third_party','workspace_provider_incidents','外部办公服务提供商事件公告')]:
    d=cloud[cloud.group==group].drop_duplicates('incident_id').copy()
    d['quarter']=pd.PeriodIndex(d.quarter,freq='Q')
    add(feature,family,label,d.groupby('quarter').size().reindex(complete_cloud,fill_value=0),
        'Google official product incident histories','unique incidents',limit='选定产品的公开服务事件，可含中断、性能下降等，数量不区分严重程度或影响范围；公告政策可能变化。未验证DTCC使用这些服务，不能当作DTCC事件。GCP技术与数据两组可以重叠。',primary=True)

# BLS rates and economic uncertainty: require all three monthly observations.
for sid,feature,family,label,units,limit in [
 ('JTU5200QUR','finance_sector_quit_rate','human_capital','金融保险行业主动离职率','percent','行业代理，非DTCC员工数据；未经季调，分析另控制季度季节性。'),
 ('USEPUINDXM','economic_policy_uncertainty','strategic','美国经济政策不确定性','index','外部战略环境代理，不证明战略项目失败或竞争地位下降。'),
 ('DGS3MO','treasury_bill_yield_volatility','corporate_financial','三个月国债收益率波动','percentage points','公司现金再投资环境代理，不测量DTCC实际投资、资本或资金缺口。')]:
    path=RAW/(sid+'.csv')
    if not path.exists():continue
    d=pd.read_csv(path);d.columns=['date','value'];d['date']=pd.to_datetime(d.date);d['value']=pd.to_numeric(d.value,errors='coerce')
    d['quarter']=d.date.dt.to_period('Q')
    d.to_csv(DATA/(sid+'_observations.csv'),index=False)
    if sid=='DGS3MO':
        series=d.groupby('quarter').value.std();counts=d.groupby('quarter').value.count();series[counts<45]=np.nan
    else:
        series=d.groupby('quarter').value.mean();counts=d.groupby('quarter').value.count();series[counts<3]=np.nan
    add(feature,family,label,series,'FRED '+sid,units,limit=limit,primary=True)

# Actual financial-regulator rule publications. Page count and IDs reconcile.
fr_rows=[];declared_counts=[]
for path in sorted(RAW.glob('federal_register_rules_*.json')):
    o=json.loads(path.read_text());fr_rows.extend(o['results']);declared_counts.append(o['count'])
fr=pd.DataFrame(fr_rows).drop_duplicates('document_number')
assert len(fr)==declared_counts[0]
fr['quarter']=pd.to_datetime(fr.publication_date).dt.to_period('Q')
fr['agency_names']=fr.agencies.map(lambda x:'; '.join(a.get('name','') for a in x))
fr.drop(columns=['agencies']).to_csv(DATA/'financial_regulator_final_rules.csv',index=False)
add('financial_regulator_final_rules','legal_regulatory','SEC/CFTC/美联储最终规则发布数量',fr.groupby('quarter').size().reindex(QUARTERS,fill_value=0),
    'Federal Register API: SEC/CFTC/Federal Reserve RULE publications','documents',limit='包含行政与其他金融业务规则，数量只测量监管变化活动；不证明FICC/ECS合规失败。',primary=True)

# Extract more of the already-downloaded DTCC disclosures, preserving cell evidence.
codes={'6.5.1.1','6.5.2','6.5.3','15.1.1','15.1.2','17.1.1','17.2.1','17.3.1'}
entities={'FICC/GSD','FICC/MBSD','NSCC','FICC','DTC','DTCC Clearing Agencies'}
pqd_rows=[]
def number(x):
    try:return float(str(x).replace(',','').replace('%','').strip())
    except (ValueError,TypeError):return np.nan
pqd_folder=RAW/'pqd_sources'
if not pqd_folder.exists():
    pqd_folder.mkdir()
    for original in (OLD/'raw/business_documents').rglob('*.xlsx'):
        (pqd_folder/original.name).write_bytes(original.read_bytes())
for path in sorted(pqd_folder.glob('*.xlsx')):
    match=re.search(r'(20\d\d)_Q([1-4])',path.name)
    if not match:continue
    quarter=pd.Period(match[1]+'Q'+match[2],freq='Q')
    w=openpyxl.load_workbook(path,data_only=True,read_only=True)
    if 'DTCC_ConsolidatedDataFile' not in w.sheetnames:continue
    for row_num,r in enumerate(w['DTCC_ConsolidatedDataFile'].iter_rows(values_only=True),1):
        code=str(r[1])
        if code not in codes:continue
        for i,v in enumerate(r[:-5]):
            if v not in entities:continue
            raw=r[i+4];note=str(r[i+5] or '')
            pqd_rows.append({'quarter':str(quarter),'disclosure':code,'entity':v,'raw_value':str(raw),'numeric_value':number(raw),
                             'note':note,'file':str(path.relative_to(ROOT)),'sheet':'DTCC_ConsolidatedDataFile',
                             'row':row_num,'value_column':i+5})
    w.close()
pqd=pd.DataFrame(pqd_rows)
assert not pqd.duplicated(['quarter','disclosure','entity']).any()
pqd.to_csv(DATA/'pqd_risk_cell_evidence.csv',index=False)
coverage_checks=[]
for entity in ['FICC/GSD','FICC/MBSD','NSCC']:
    b=pqd[pqd.entity==entity].pivot(index='quarter',columns='disclosure',values='numeric_value')
    b.index=pd.PeriodIndex(b.index,freq='Q')
    c=b['6.5.3'].copy();c.loc[c<=1]*=100
    expected=100*(1-b['6.5.1.1']/b['6.5.2'])
    for q in b.index:
        coverage_checks.append({'quarter':str(q),'entity':entity,'reported_coverage_pct':c[q],'count_implied_coverage_pct':expected[q],'difference_pp':c[q]-expected[q]})
    model=100-c
    add('margin_backtest_uncovered_pct','model_risk','保证金模型回测未覆盖比例',model,'DTCC PQD 6.5.3','percent',
        scope='own_disclosed_metric',entity=entity,window='trailing 12 months',limit='过去12个月回测覆盖结果；与模型使用环境及市场波动有关，不等于设计或实现错误，也不是已实现损失。',primary=True)
    availability=b['17.2.1'].copy();availability.loc[availability<=1]*=100
    add('core_system_unavailability_pct','processing_operations','核心系统不可用比例',100-availability,'DTCC PQD 17.2.1','percent',
        scope='own_disclosed_metric',entity=entity,window='specified availability window / system aggregate',limit='系统聚合可用性口径，部分报告涉及共享系统；不能识别事故根因或直接归属每个业务。',primary=False)
for entity in ['FICC','NSCC']:
    d=pqd[(pqd.entity==entity)&pqd.disclosure.isin(['15.1.1','15.1.2'])].copy()
    # Old reports repeat annual statements; only explicitly quarter-end rows qualify.
    d=d[d.note.str.contains('Quarter-end',case=False,na=False)]
    b=d.pivot(index='quarter',columns='disclosure',values='numeric_value');b.index=pd.PeriodIndex(b.index,freq='Q')
    ratio=b['15.1.2']/b['15.1.1']
    add('own_liquidity_buffer_pressure','corporate_financial','半年运营费用/自有流动净资产',ratio,'DTCC PQD 15.1.1/15.1.2','ratio',
        scope='own_disclosed_metric',entity=entity,window='quarter end',limit='比值越高缓冲相对越小；费用与财务表现机械相关，属于内生状态指标；早期年度重复数据已排除。')
shared=pqd[(pqd.disclosure=='17.3.1') & (pqd.entity=='DTCC Clearing Agencies')].copy()
counts=shared.raw_value.str.extract(r'^\s*(\d+)\s*/')[0].astype(float)
series=pd.Series(counts.values,index=pd.PeriodIndex(shared.quarter,freq='Q'))
add('clearing_core_failures_12m','processing_operations','清算核心系统过去12个月故障数',series,'DTCC PQD 17.3.1','failures',
    scope='shared_clearing_agency_metric',entity='CLEARING',window='trailing 12 months',limit='GSD/MBSD/NSCC共享清算系统口径；不能当作单季新增事件，不能分配给DTC或某单一业务。',primary=True)
pd.DataFrame(coverage_checks).to_csv(DATA/'backtest_coverage_reconciliation.csv',index=False)

# FICC-level backtest summary weights the two division coverages by observations.
bt=pqd[pqd.entity.isin(['FICC/GSD','FICC/MBSD'])].pivot(index=['quarter','entity'],columns='disclosure',values='numeric_value')
bt['coverage_pct']=bt['6.5.3'].where(bt['6.5.3']>1,bt['6.5.3']*100)
weighted={}
for q,d in bt.groupby(level='quarter'):
    if len(d)==2 and d[['coverage_pct','6.5.2']].notna().all().all():
        weighted[pd.Period(q,freq='Q')]=float(np.average(100-d.coverage_pct,weights=d['6.5.2']))
add('margin_backtest_uncovered_pct','model_risk','保证金模型回测未覆盖比例：FICC汇总',pd.Series(weighted),
    'DTCC PQD 6.5.3 with 6.5.2 observation weights','percent',scope='own_disclosed_metric',entity='FICC',
    window='trailing 12 months',limit='按GSD与MBSD回测观察数加权的研究汇总，非DTCC发布的独立FICC风险评级；不是已实现损失。',primary=True)

defs=pd.DataFrame(definitions)
assert set(defs.family)==set(x[0] for x in families)
defs.to_csv(DATA/'expanded_risk_dictionary.csv',index=False)
observations=pd.DataFrame(records)
assert not observations.duplicated(['quarter','feature','entity_scope']).any()
observations.to_csv(DATA/'expanded_risk_observations.csv',index=False)
coverage=observations.groupby(['feature','entity_scope']).agg(first_quarter=('quarter','min'),last_quarter=('quarter','max'),quarters=('value','count'),unique_values=('value','nunique')).reset_index()
coverage=coverage.merge(defs,on=['feature','entity_scope'],validate='one_to_one')
coverage.to_csv(DATA/'collection_coverage.csv',index=False)
print('Risk families with data:',len(set(defs.family)),'features:',defs.feature.nunique(),'series:',len(defs),'observations:',len(observations))
print(coverage[['family','feature','entity_scope','first_quarter','last_quarter','quarters','unique_values']].to_string(index=False))
