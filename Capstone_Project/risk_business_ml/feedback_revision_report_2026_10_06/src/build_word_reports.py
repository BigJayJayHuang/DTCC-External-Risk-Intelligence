from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
C = json.loads((ROOT / 'report_content.json').read_text())
OUT = ROOT / 'output'
OUT.mkdir(exist_ok=True)

def font(run, size=10.5, bold=False):
    run.font.name = 'Arial'
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Arial Unicode MS')

def para(doc, text, style=None, size=None):
    p = doc.add_paragraph(style=style)
    r = p.add_run(text)
    font(r, size or (10.5 if style is None else 10))
    return p

def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    font(p.add_run(text), 15 if level == 1 else 12, True)
    return p

def table(doc, headers, rows, widths, numeric=()):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths): col.width = Inches(width)
    for i,h in enumerate(headers): t.rows[0].cells[i].text = h
    for values in rows:
        cells = t.add_row().cells
        for i,v in enumerate(values): cells[i].text = str(v)
    for ri,row in enumerate(t.rows):
        trPr = row._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit'); trPr.append(no_split)
        if ri == 0:
            repeat = OxmlElement('w:tblHeader'); trPr.append(repeat)
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
            pr.append(borders)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side); e.set(qn('w:w'),'85' if side in ['top','bottom'] else '95'); e.set(qn('w:type'),'dxa'); margins.append(e)
            pr.append(margins)
            if ri==0:
                shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E8EDF1'); pr.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(0)
                p.paragraph_format.space_after=Pt(0)
                p.paragraph_format.line_spacing=1.12
                if ci in numeric: p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
                for r in p.runs:font(r,9,ri==0)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def picture(doc, name, caption, alt):
    p=doc.add_paragraph(); p.paragraph_format.keep_with_next=True
    r=p.add_run(); shape=r.add_picture(str(ROOT/'figures'/name),width=Inches(6.65))
    shape._inline.docPr.set('descr',alt)
    p=para(doc,caption,size=9)
    p.paragraph_format.space_after=Pt(12)

def hyperlink(p,label,url):
    h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r'); pr=OxmlElement('w:rPr')
    color=OxmlElement('w:color'); color.set(qn('w:val'),'145E75');pr.append(color)
    sz=OxmlElement('w:sz');sz.set(qn('w:val'),'19');pr.append(sz)
    r.append(pr);text=OxmlElement('w:t');text.text=label;r.append(text);h.append(r);p._p.append(h)

def build(lang):
    cn=lang=='cn'; idx=0 if cn else 1; copy=C[lang]
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
    sec.top_margin=sec.bottom_margin=Inches(.72)
    sec.left_margin=sec.right_margin=Inches(.8)
    sec.footer_distance=Inches(.32)
    normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10.5)
    normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Arial Unicode MS')
    normal.paragraph_format.line_spacing=1.2
    normal.paragraph_format.space_after=Pt(8)
    for st in doc.styles:
        for border in list(st._element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    for name in ['Heading 1','Heading 2']:
        doc.styles[name].paragraph_format.keep_with_next=True
        doc.styles[name].paragraph_format.space_before=Pt(12)
        doc.styles[name].paragraph_format.space_after=Pt(7)
    doc.core_properties.title=copy['title']; doc.core_properties.subject=copy['subtitle']
    doc.core_properties.author='Capstone Project'
    p=doc.add_paragraph(style='Title');font(p.add_run(copy['title']),23,True)
    p.paragraph_format.space_after=Pt(10)
    para(doc,copy['subtitle'],size=11)
    para(doc,'2026年10月6日   数据观察截至2026Q2' if cn else 'October 6, 2026   Observations through 2026Q2',size=9.5)
    para(doc,'事后扩展的探索性研究' if cn else 'Post-result exploratory extension',size=9.5)
    for si,section in enumerate(copy['sections']):
        if si in [1,2,3,4]:doc.add_page_break()
        heading(doc,section['heading'])
        for text in section['paragraphs']:para(doc,text)
        if si==2:
            headers=['指标','公式','窗口及单位','解释限制'] if cn else ['Measure','Formula','Window and unit','Interpretation limit']
            rows=[[C['outcome_labels'][f[0]][idx],f[1],f[2 if cn else 3],f[4 if cn else 5]] for f in C['fields']]
            table(doc,headers,rows,[1.45,1.6,1.2,2.42])
        if si==3:
            picture(doc,'period_and_lag_comparison.png',
                '图1 不同时期与滞后下的相关性  后续期只有六对观察' if cn else 'Figure 1 Associations by period and lag  The later sample has only six pairs',
                'Earlier and later Spearman correlations for GSD and MBSD Treasury changes versus peak VM, and NSCC VIX versus peak VM.')
            case_heading=heading(doc,'展示案例的完整统计' if cn else 'Complete statistics for displayed cases',level=2)
            case_heading.paragraph_format.page_break_before=True
            case_keys=[]
            for ci,r in enumerate(C['cases'][::2]):
                key=chr(65+ci)
                desc=f"{key}  {r['entity']}  {C['indicator_labels'][r['indicator']][idx]} → {C['outcome_labels'][r['outcome']][idx]}"
                case_keys.append(desc)
            para(doc,'\n'.join(case_keys),size=9.5)
            headers=['案例','滞后','n','rho','BH q','95%区间','较早rho','后续rho'] if cn else ['Case','Lag','n','rho','BH q','95% interval','Earlier rho','Later rho']
            rows=[]
            for ci,r in enumerate(C['cases']):
                rows.append([chr(65+ci//2),r['lag'],r['n'],f"{r['rho']:.3f}",f"{r['q']:.3f}",f"{r['ci95_low']:.3f} to {r['ci95_high']:.3f}",f"{r['rho_earlier']:.3f}",f"{r['rho_later']:.3f}"])
            table(doc,headers,rows,[.48,.48,.35,.62,.62,1.58,1.06,1.08],numeric=[1,2,3,4,6,7])
            para(doc,'滞后0为同季度，1为下一季度。区间来自全样本四季度移动区块bootstrap。较早结果季度截至2024Q4，后续为2025Q1至2026Q2。q在全样本294项比较中校正。' if cn else 'Lag 0 is same-quarter; lag 1 is next-quarter. Intervals use the full-sample four-quarter moving-block bootstrap. Earlier target quarters end in 2024Q4; later targets are 2025Q1–2026Q2. q-values adjust the full family of 294 comparisons.',size=9)
        if si==4:
            picture(doc,'same_quarter_cases.png',
                '图2 全样本同季度相关性与95%区间  虚线为原研究的0.3方向门槛  不代表因果或验证通过' if cn else 'Figure 2 Full-sample same-quarter associations and 95% intervals  The dashed line marks the original 0.3 directional threshold  It does not establish causality or validation',
                'Six predefined same-quarter correlations with four-quarter block-bootstrap confidence intervals; wide intervals and mixed signs show uncertainty.')
        if si==6:
            for source in C['sources']:
                p=doc.add_paragraph();p.paragraph_format.space_after=Pt(5)
                hyperlink(p,source['label'],source['url'])
            heading(doc,'可复现文件' if cn else 'Reproducible files',level=2)
            para(doc,'分析程序 src/analyze_responses.py\n原始字段 data/response_cell_evidence.csv\n全部结果 results/response_correlations.csv\n来源记录 source_manifest.json' if cn else 'Analysis script src/analyze_responses.py\nRaw field evidence data/response_cell_evidence.csv\nComplete results results/response_correlations.csv\nSource manifest source_manifest.json',size=9)
    f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    font(f.add_run('Capstone  |  '),8)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
    path=OUT/f'PQD_Feedback_Review_and_Supplemental_Tests_{lang.upper()}.docx'
    doc.save(path)
    print(path)

for lang in ['cn','en']:build(lang)
