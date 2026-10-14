"""Exploratory response diagnostics added after review of the original results.
Run with risk_business_ml/.venv/bin/python. Never overwrites the original analysis.
"""
from pathlib import Path
import json, hashlib, datetime, importlib.util, warnings
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.utils import get_column_letter
from scipy.stats import spearmanr, rankdata
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

ROOT=Path(__file__).resolve().parents[1]
OLD=ROOT.parent/'event_vulnerability_tests_2026_10_06'
ENT={'FICC/GSD':'GSD','FICC/MBSD':'MBSD','NSCC':'NSCC'}
OUTCOMES={
 'im_qoq':('Quarter-end IM growth','fraction','IM(t)/IM(t-1)-1'),
 'im_call_usd':('Peak daily IM call','USD','PQD 6.8.1'),
 'im_call_to_im':('Peak IM call / quarter-end IM','fraction','6.8.1(t)/6.1.1(t)'),
 'im_call_to_prior_im':('Peak IM call / prior-quarter IM','fraction','6.8.1(t)/6.1.1(t-1); sensitivity'),
 'vm_avg_usd':('Average daily VM','USD','PQD 6.6.1'),
 'vm_peak_usd':('Peak daily VM','USD','PQD 6.7.1'),
 'vm_peak_to_avg':('Peak / average VM','multiple','6.7.1(t)/6.6.1(t)'),
}
LABEL={'sofr_effr':'SOFR-EFFR','sofr_effr_adjusted':'Adjusted SOFR-EFFR','ust10_abs_change':'10Y absolute change','ust10_rv20':'10Y yield RV20','baa_spread':'Baa spread','baa_widen20':'Baa widening 20D','vix':'VIX'}
plan={'version':'response-review-v1','prepared_at_utc':datetime.datetime.now(datetime.timezone.utc).isoformat(),'scope':['GSD','MBSD','NSCC'],'cutoff':'2026Q2','status':'exploratory revision after original outcomes and later-sample results were seen','external_indicators':'reuse original seven transformations and quarterly maxima; do not change event windows or event results','outcomes':OUTCOMES,'lags':{'0':'same-quarter association','1':'next-quarter association; not a deployable forecast'},'samples':{'earlier':'target quarters through 2024Q4, no longer a fresh selection test','later':'2025Q1 through 2026Q2, already observed; only six quarters','all':'2020Q4 through 2026Q2'},'inference':'ordinary Spearman p, BH within all 294 comparisons of each sample; four-consecutive-quarter moving block bootstrap 2000 replicates for n>=12; calendar-respecting first differences','no_untouched_final_test':True,'no_causal_or_calibrated_score':True,'predefined_display_cases':[['GSD','ust10_abs_change','im_qoq'],['GSD','ust10_abs_change','im_call_to_im'],['GSD','ust10_abs_change','vm_peak_usd'],['MBSD','ust10_abs_change','vm_peak_usd'],['NSCC','vix','vm_peak_usd'],['GSD','sofr_effr_adjusted','vm_peak_to_avg']]}
(ROOT/'analysis_plan.json').write_text(json.dumps(plan,ensure_ascii=False,indent=2))
# Hash relevant original results before/after to prove they are preserved.
original_files=[OLD/'analysis_plan.json',OLD/'data/pqd_quarterly.csv',OLD/'results/event_screen.csv',OLD/'results/pqd_lagged_spearman.csv',OLD/'results/combined_screen_and_research_weights.csv']
original_hash={str(p.relative_to(ROOT.parent)):hashlib.sha256(p.read_bytes()).hexdigest() for p in original_files}
manifest=json.loads((OLD/'source_manifest.json').read_text()); evid=[]; sources=[]
def number(x):
 if x is None or str(x).strip().lower() in ['', 'n/a','na','none','-']:return np.nan
 return float(str(x).replace(',',''))
for m in manifest:
 p=OLD/m['local_file']
 if p.suffix not in ['.xlsx','.pdf'] or '/pqd/' not in m['local_file']:continue
 assert hashlib.sha256(p.read_bytes()).hexdigest()==m['sha256'],p.name
 sources.append({**m,'reused_local_file':str(p.relative_to(ROOT.parent)), 'reused_for':'business response diagnostics'})
 if p.suffix!='.xlsx':continue
 q=f"{m['reporting_year']}Q{m['reporting_quarter']}"
 w=openpyxl.load_workbook(p,read_only=True,data_only=True)
 s=w['DTCC_ConsolidatedDataFile']
 for nr,r in enumerate(s.values,1):
  if len(r)<2 or str(r[1]) not in ['6.6.1','6.7.1','6.8.1']:continue
  code=str(r[1]);assert pd.Timestamp(r[0]).to_period('Q')==pd.Period(q,'Q')
  for i,x in enumerate(r):
   if x in ENT:
    evid.append({'quarter':q,'entity':ENT[x],'code':code,'value_usd':number(r[i+4]),'raw_value':str(r[i+4]),'raw_unit':'USD','scale_to_usd':1,'description':r[3],'note':r[i+5],'source_file':str(p.relative_to(ROOT.parent)),'source_url':m['source_urls'][0],'sheet':s.title,'cell':f'{get_column_letter(i+5)}{nr}','page':None,'source_sha256':m['sha256']})
 w.close()
# Official 2024Q4 PDF p13, amounts explicitly USD MM. Cross-check saved page text.
pdf=next(m for m in manifest if m['local_file'].endswith('DTCC_PQD_2024_Q4.pdf'))
page_text=(OLD/'data/2024Q4_page_13.txt').read_text()
for code,vals in [('6.6.1',[10679,548,557]),('6.7.1',[14505,1817,3954]),('6.8.1',[7574,2189,7230])]:
 assert code in page_text
 for e,v in zip(['GSD','MBSD','NSCC'],vals):
  assert f'{v:,}' in page_text
  evid.append({'quarter':'2024Q4','entity':e,'code':code,'value_usd':v*1e6,'raw_value':v,'raw_unit':'USD MM','scale_to_usd':1e6,'description':'Official PDF quarterly margin call disclosure','note':'6.8.1 actual aggregate deficit over deposits; includes intraday calls. GSD/MBSD VM cash pass-through; NSCC VM part of clearing fund requirements.','source_file':str((OLD/pdf['local_file']).relative_to(ROOT.parent)),'source_url':pdf['source_urls'][0],'sheet':'','cell':f'{e} column for {code}','page':13,'source_sha256':pdf['sha256']})
e=pd.DataFrame(evid);assert len(e)==207 and not e.duplicated(['quarter','entity','code']).any();e.to_csv(ROOT/'data/response_cell_evidence.csv',index=False,encoding='utf-8-sig')
panel=pd.read_csv(OLD/'data/pqd_quarterly.csv')
wide=e.pivot(index=['quarter','entity'],columns='code',values='value_usd').reset_index().rename(columns={'6.6.1':'vm_avg_usd','6.7.1':'vm_peak_usd','6.8.1':'im_call_usd'})
panel=panel.merge(wide,on=['quarter','entity'],validate='one_to_one').sort_values(['entity','quarter'])
assert len(panel)==69
panel['prior_im']=panel.groupby('entity').initial_margin_required_usd.shift()
panel['im_qoq']=panel.initial_margin_required_usd/panel.prior_im-1
panel['im_call_to_im']=panel.im_call_usd/panel.initial_margin_required_usd
panel['im_call_to_prior_im']=panel.im_call_usd/panel.prior_im
panel['vm_peak_to_avg']=panel.vm_peak_usd/panel.vm_avg_usd
assert (panel.vm_peak_usd>=panel.vm_avg_usd).all() and (panel.im_call_usd>=0).all()
assert all((panel.groupby('entity').size()==23))
panel.to_csv(ROOT/'data/response_quarterly.csv',index=False,encoding='utf-8-sig')
q=pd.read_csv(OLD/'data/quarterly_external_pressure.csv');ev=pd.read_csv(OLD/'results/event_screen.csv')
def rho(x,y):
 if len(x)<3 or np.unique(x).size<2 or np.unique(y).size<2:return np.nan
 return float(spearmanr(x,y).statistic)
def ci(x,y,dates,seed):
 n=len(x);b=4;k=int(np.ceil(n/b));o=np.array([pd.Period(s,'Q').ordinal for s in dates]);starts=[i for i in range(n-b+1) if np.all(np.diff(o[i:i+b])==1)]
 if not starts:return np.nan,np.nan
 ids=(np.random.default_rng(seed).choice(starts,size=(2000,k))[:,:,None]+np.arange(b)).reshape(2000,-1)[:,:n]
 xr=rankdata(x[ids],axis=1);yr=rankdata(y[ids],axis=1);xr-=xr.mean(1,keepdims=True);yr-=yr.mean(1,keepdims=True)
 den=np.sqrt((xr*xr).sum(1)*(yr*yr).sum(1));vals=np.divide((xr*yr).sum(1),den,out=np.full(2000,np.nan),where=den>0)
 return np.nanquantile(vals,[.025,.975])
def bh(a):
 v=np.asarray(a,float);out=np.full(len(v),np.nan);valid=np.flatnonzero(np.isfinite(v));o=valid[np.argsort(v[valid])];n=len(o)
 if n:out[o]=np.minimum(1,np.minimum.accumulate((v[o]*n/np.arange(1,n+1))[::-1])[::-1])
 return out
rows=[];pairs=[]
for ind,x0 in q.groupby('indicator'):
 for lag in [0,1]:
  x=x0.copy();x['target_quarter']=x.quarter.map(lambda s:str(pd.Period(s,'Q')+lag))
  for entity,y in panel.groupby('entity'):
   pair=x.merge(y.rename(columns={'quarter':'target_quarter'}),on='target_quarter',validate='one_to_one').sort_values('target_quarter')
   for outcome in OUTCOMES:
    pairs.append(pair[['quarter','target_quarter','indicator','stress','entity',outcome]].rename(columns={'quarter':'stress_quarter',outcome:'outcome_value'}).assign(outcome=outcome,lag=lag))
    for sample,mask in [('earlier',pair.target_quarter<='2024Q4'),('later',pair.target_quarter>'2024Q4'),('all',pd.Series(True,index=pair.index))]:
     part=pair[mask].dropna(subset=['stress',outcome]);xa=part.stress.to_numpy();ya=part[outcome].to_numpy();n=len(part)
     with warnings.catch_warnings():
      warnings.simplefilter('ignore');st=spearmanr(xa,ya) if n>=3 and np.unique(xa).size>1 and np.unique(ya).size>1 else None
     r=float(st.statistic) if st is not None else np.nan;p=float(st.pvalue) if st is not None else np.nan
     sufficient=n>=12 and np.unique(xa).size>=3 and np.unique(ya).size>=3
     low,high=ci(xa,ya,part.target_quarter,2026100600+len(rows)) if sufficient else [np.nan,np.nan]
     lev=part.set_index('target_quarter')[['stress',outcome]];prev=lev.copy();prev.index=[str(pd.Period(s,'Q')+1) for s in prev.index];diff=(lev-prev).dropna()
     rows.append({'entity':entity,'indicator':ind,'channel':x0.channel.iloc[0],'economic_role':'auxiliary' if (x0.channel.iloc[0]=='equity' and entity!='NSCC') or (x0.channel.iloc[0]=='rates' and entity=='NSCC') else 'primary','outcome':outcome,'lag':lag,'sample':sample,'n':n,'rho':r,'p_ordinary':p,'ci95_low':low,'ci95_high':high,'rho_first_differences':rho(diff.stress.values,diff[outcome].values),'sufficient_n':sufficient,'positive_size_flag':bool(sufficient and r>=.3),'event_gate_pass':bool(ev.set_index('indicator').loc[ind,'event_gate_pass']),'status':'EXPLORATORY_POST_RESULT'})
res=pd.DataFrame(rows);res['bh_q_ordinary']=np.nan
for sample,g in res.groupby('sample'):res.loc[g.index,'bh_q_ordinary']=bh(g.p_ordinary)
res['ordinary_fdr_5pct']=res.bh_q_ordinary<=.05
res['positive_ci_and_ordinary_fdr']=res.ordinary_fdr_5pct & res.ci95_low.gt(0)
res.to_csv(ROOT/'results/response_correlations.csv',index=False,encoding='utf-8-sig')
pd.concat(pairs,ignore_index=True).to_csv(ROOT/'data/response_pair_evidence.csv',index=False,encoding='utf-8-sig')
case=[]
for entity,ind,outcome in plan['predefined_display_cases']:
 for lag in [0,1]:
  g=res[(res.entity==entity)&(res.indicator==ind)&(res.outcome==outcome)&(res.lag==lag)]
  a=g.set_index('sample');full=a.loc['all'];case.append({'entity':entity,'indicator':ind,'outcome':outcome,'lag':lag,'n':int(full.n),'rho':float(full.rho),'ci95_low':float(full.ci95_low),'ci95_high':float(full.ci95_high),'q':float(full.bh_q_ordinary),'rho_earlier':float(a.loc['earlier','rho']),'rho_later':float(a.loc['later','rho']),'rho_diff':float(full.rho_first_differences)})
case=pd.DataFrame(case);case.to_csv(ROOT/'results/display_cases.csv',index=False,encoding='utf-8-sig')
# Figures retain EN labels; both reports use identical figures with translated captions.
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10,'axes.spines.top':False,'axes.spines.right':False})
fig,ax=plt.subplots(figsize=(9.2,5.2));sub=case[case.lag==0].iloc[::-1]
labels=[f"{r.entity}: {LABEL[r.indicator]}\n{OUTCOMES[r.outcome][0]}" for r in sub.itertuples()]
y=np.arange(len(sub));ax.axvline(0,color='#a0a7b0',lw=1);ax.axvline(.3,color='#b49a60',ls='--',lw=.8)
for yi,r in zip(y,sub.itertuples()):ax.plot([r.ci95_low,r.ci95_high],[yi,yi],color='#778ca0',lw=2);ax.scatter(r.rho,yi,c='#176a69',s=50,zorder=3)
ax.set_yticks(y,labels);ax.set_xlim(-1,1);ax.set_xlabel('Spearman rho with four-quarter block bootstrap 95% interval');ax.set_title('Same-quarter business responses: predefined descriptive examples',loc='left',fontweight='bold');fig.tight_layout();fig.savefig(ROOT/'figures/same_quarter_cases.png',dpi=180);plt.close(fig)
fig,axs=plt.subplots(1,3,figsize=(10,3.5),sharey=True)
for ax,(entity,ind,outcome) in zip(axs,[plan['predefined_display_cases'][2],plan['predefined_display_cases'][3],plan['predefined_display_cases'][4]]):
 part=case[(case.entity==entity)&(case.indicator==ind)&(case.outcome==outcome)]
 for idx,samp in enumerate(['rho_earlier','rho_later']):ax.bar(np.arange(2)+(-.17 if idx==0 else .17),part.sort_values('lag')[samp],width=.32,label='Through 2024Q4' if idx==0 else '2025Q1-2026Q2',color=['#176a69','#dca164'][idx])
 ax.axhline(0,color='#a0a7b0',lw=.8);ax.set_xticks([0,1],['Same quarter','Next quarter']);ax.set_ylim(-1,1);ax.set_title(f'{entity}\n{LABEL[ind]} vs peak VM')
axs[0].set_ylabel('Spearman rho');axs[-1].legend(fontsize=8,loc='lower right');fig.suptitle('Relationships can change across periods; later sample has only six pairs',fontsize=12);fig.tight_layout();fig.savefig(ROOT/'figures/period_and_lag_comparison.png',dpi=180);plt.close(fig)
# Provenance includes frozen source files, definitions, and reused evidence hashes.
source_record={'generated_at_utc':datetime.datetime.now(datetime.timezone.utc).isoformat(),'scope':'new fields extracted from previously collected official files; no new historical PQD collected','sources':sources,'reused_derived_inputs':[{ 'file':str(p.relative_to(ROOT.parent)),'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in [OLD/'data/pqd_quarterly.csv',OLD/'data/quarterly_external_pressure.csv',OLD/'results/event_screen.csv']], 'definition_source':'https://www.bis.org/publications/public-quantitative-disclosure-standards-central-counterparties.pdf','pdf_values':'2024Q4 p13, USD MM scaled by 1000000; original page text retained in original analysis'}
(ROOT/'source_manifest.json').write_text(json.dumps(source_record,ensure_ascii=False,indent=2))
assert all(hashlib.sha256((ROOT.parent/n).read_bytes()).hexdigest()==h for n,h in original_hash.items())
counts={sample:{'comparisons':len(g),'positive_rho_ge_03_n_ge_12':int(g.positive_size_flag.sum()),'ordinary_fdr_5pct':int(g.ordinary_fdr_5pct.sum()),'positive_ci_and_ordinary_fdr':int(g.positive_ci_and_ordinary_fdr.sum())} for sample,g in res.groupby('sample')}
summary={'date':'2026-10-06','cutoff':'2026Q2','quarters':23,'entities':3,'new_field_cells':207,'external_indicators':7,'response_formulations':7,'comparisons':len(res),'counts':counts,'original_results_preserved':True,'original_result_hashes':original_hash,'exploratory_only':True,'cases':case.replace({np.nan:None}).to_dict('records')}
(ROOT/'results/analysis_summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2))
print(json.dumps({k:v for k,v in summary.items() if k not in ['cases','original_result_hashes']},ensure_ascii=False))
print(case.round(3).to_string(index=False))
