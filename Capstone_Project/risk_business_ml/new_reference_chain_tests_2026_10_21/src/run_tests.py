"""Offline, reproducible tests of new v3 routes. No proxy result validates a full chain."""
from pathlib import Path
from datetime import datetime,timezone
from zipfile import ZipFile
import io,json,re,hashlib,warnings
import numpy as np
import pandas as pd
from scipy.stats import spearmanr,rankdata
import statsmodels.api as sm
from statsmodels.stats.multitest import multipletests
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
PLAN=json.loads((ROOT/'analysis_plan.json').read_text());SEED=20261016
def save(frame,name,sub='data'):
    frame.to_csv(ROOT/sub/name,index=False)
def dump(value,name):
    (ROOT/name).write_text(json.dumps(value,ensure_ascii=False,indent=2,default=str,allow_nan=False)+'\n')
def clean(value):
    if isinstance(value,dict):return {k:clean(v) for k,v in value.items()}
    if isinstance(value,list):return [clean(v) for v in value]
    if isinstance(value,(float,np.floating)) and not np.isfinite(value):return None
    if isinstance(value,np.generic):return value.item()
    return value
def check_inputs():
    for row in json.loads((ROOT/'input_manifest.json').read_text()):
        assert hashlib.sha256((ROOT/row['frozen_copy']).read_bytes()).hexdigest()==row['sha256']
def numeric(value):
    try:return float(str(value).replace(',','').replace('%',''))
    except (TypeError,ValueError):return np.nan
def q_end(d):return d.weekday()<5 and (d+pd.offsets.BDay(3)).quarter!=d.quarter
def rho(x,y):
    if len(x)<3 or np.unique(x).size<2 or np.unique(y).size<2:return np.nan
    return float(spearmanr(x,y).statistic)

def extract_pqd():
    rows=[]
    refs={'18.2.2','18.3.2','20.7.2','23.1.2','23.2.1'}
    for path in sorted((ROOT/'inputs/pqd').glob('*.xlsx')):
        book=openpyxl.load_workbook(path,data_only=True,read_only=True)
        sheet=book['DTCC_ConsolidatedDataFile'];headers=list(next(sheet.iter_rows(max_row=1,values_only=True)))
        for rn,row in enumerate(sheet.iter_rows(min_row=2,values_only=True),2):
            ref=str(row[1])
            if ref not in refs:continue
            for entity,g in [('GSD',1),('MBSD',2),('NSCC',3)]:
                col=headers.index('Value'+str(g));description=str(row[col-1] or '')
                raw=row[col];v=numeric(raw)
                if ref in ['18.2.2','18.3.2']:
                    # Match prior code: strings marked % or numeric percentages >1 become fractions.
                    if '%' in str(raw) or v>1:v/=100
                    unit='fraction';normalization='divide by 100 if explicit percent sign or numeric >1; otherwise retain source fraction'
                elif ref=='20.7.2':
                    unit='percentage points as reported';normalization='retain reported number; used descriptively only'
                else:
                    unit='USD as reported; reconcile23.2.1 before use' if ref=='23.2.1' else 'USD'
                    normalization='unscaled ValueN numeric field, not formatted display columns;23.2.1 requires financial-note reconciliation'
                rows.append(dict(quarter=path.stem.replace('DTCC_PQD_','').replace('_',''),entity=entity,reference=ref,
                  dimension=description,definition=row[3],raw_value=raw,value=v,unit=unit,normalization=normalization,
                  reported_entity=row[col-4],comments=row[col+1],source_file=path.name,sheet=sheet.title,
                  cell=f'{openpyxl.utils.get_column_letter(col+1)}{rn}',sha256=hashlib.sha256(path.read_bytes()).hexdigest()))
        book.close()
    frame=pd.DataFrame(rows);save(frame,'new_pqd_cell_evidence.csv')
    selected=[]
    for ref,field in [('18.2.2','top5_position_share'),('18.3.2','top5_im_share'),('20.7.2','crossmargin_reduction_reported_pct'),('23.1.2','cleared_notional_usd_as_reported'),('23.2.1','unsettled_notional_usd_as_reported')]:
        a=frame[frame.reference.eq(ref)].copy()
        if ref in ['18.2.2','18.3.2']:a=a[a.dimension.str.contains('Average',case=False)]
        if ref in ['23.1.2','23.2.1']:a=a[a.dimension.str.startswith('OTC')]
        assert not a.duplicated(['quarter','entity']).any(),(ref,a[a.duplicated(['quarter','entity'],False)].to_dict('records'))
        a=a[['quarter','entity','value']].rename(columns={'value':field});selected.append(a.set_index(['quarter','entity']))
    panel=pd.concat(selected,axis=1).reset_index();save(panel,'new_pqd_fields.csv')
    return panel

def reconcile_mbsd_activity(panel):
    """Resolve two million-scaled source cells against independent monthly totals.

    The chosen x1e6 conversion must match quarterly MBSD clearing-par totals
    divided by observed Treasury business dates within half a million USD.
    Unresolved small currency fields stay missing, never guessed from correlations.
    """
    source=pd.read_csv(ROOT/'inputs/ficc_monthly.csv')
    source['period']=pd.PeriodIndex(source.reporting_month,freq='M').asfreq('Q')
    totals=source[source.division.eq('MBSD')&source.metric_code.eq('clearingpar')].groupby('period').value_raw_source_units.sum()*1e12
    market=pd.read_csv(ROOT/'inputs/market_daily.csv')
    dates=market[market.series_id.eq('DGS10')].dropna(subset=['value']).copy()
    dates['period']=pd.to_datetime(dates.observation_date).dt.to_period('Q');days=dates.groupby('period').size()
    rows=[]
    for row in panel[panel.entity.eq('MBSD')].itertuples():
        period=pd.Period(row.quarter,freq='Q');raw=row.cleared_notional_usd_as_reported;chosen=raw;factor=1.;status='source USD as reported'
        inferred=totals.get(period,np.nan)/days.get(period,np.nan)
        if pd.notna(raw) and raw<1e8:
            if pd.notna(inferred) and abs(raw*1e6-inferred)<=501000:
                factor=1e6;chosen=raw*factor;status='million-USD encoding reconciles to independent monthly clearing-par total / Treasury observation days within rounding precision'
            else:chosen=np.nan;factor=np.nan;status='unresolved scale; excluded'
        rows.append(dict(quarter=row.quarter,raw_pqd23_1_2=raw,chosen_usd=chosen,conversion_factor=factor,monthly_total_usd=totals.get(period,np.nan),observed_treasury_days=days.get(period,np.nan),independent_daily_mean_usd=inferred,status=status))
    evidence=pd.DataFrame(rows);save(evidence,'mbsd_cleared_notional_reconciliation.csv')
    return evidence.assign(period=lambda d:pd.PeriodIndex(d.quarter,freq='Q')).set_index('period').chosen_usd.sort_index()

def extract_fund_flows():
    frames=[]
    with ZipFile(ROOT/'raw/fed_z1_20260911.zip') as z:
        for filename,col,sid in [('csv/S123_t_tu.csv','FU632051000.Q','mmf_repo_flow_million_usd'),('csv/S124_1_t_tu.csv','FU653164205.Q','mutual_fund_net_issuance_million_usd')]:
            frame=pd.read_csv(io.BytesIO(z.read(filename)))
            a=frame[['date',col]].copy();a['period']=pd.PeriodIndex(a.date.str.replace(':',''),freq='Q')
            a[sid]=pd.to_numeric(a[col],errors='coerce');a=a.set_index('period')[[sid]].loc['2000Q1':'2026Q2']
            a['quarter']=a.index.astype(str);a['source_file']=filename;a['source_series']=col
            save(a.reset_index(drop=True),sid+'.csv');frames.append(a[[sid]])
    return pd.concat(frames,axis=1)

def ofr_series(sid):
    a=pd.DataFrame(json.loads((ROOT/'raw'/f'{sid}.json').read_text()),columns=['date','value'])
    a['date']=pd.to_datetime(a.date);a['value']=pd.to_numeric(a.value,errors='coerce')
    assert not a.duplicated('date').any()
    return a.set_index('date').value.dropna().sort_index().loc[:PLAN['new_source_request_cutoff']]

def build_signals():
    market=pd.read_csv(ROOT/'inputs/market_daily.csv',parse_dates=['observation_date'])
    assert not market.duplicated(['series_id','observation_date']).any()
    wide=market.pivot(index='observation_date',columns='series_id',values='value').sort_index().loc[:PLAN['market_snapshot_cutoff']]
    series=lambda sid:wide[sid].dropna()
    basis=(pd.concat([series('SOFR'),series('EFFR')],axis=1).dropna().eval('SOFR-EFFR'))*100
    flags=np.array([q_end(d) for d in basis.index]);adj=[]
    for i,x in enumerate(basis):
        past=basis.iloc[max(0,i-252):i];f=flags[max(0,i-252):i];same=past.iloc[np.flatnonzero(f==flags[i])]
        adj.append(x-(same.median() if len(same)>=6 else past.median()) if len(past)>=126 else np.nan)
    sig={'funding':pd.Series(adj,index=basis.index).dropna(),'ust10_abs_change':series('DGS10').diff().abs()*100,'vix':series('VIXCLS')}
    cp=[]
    for sid in ['DCPF3M','DTB3']:
        f=pd.read_csv(ROOT/'inputs'/f'{sid}.csv');dc=next(c for c in f if c.lower() in ['date','observation_date']);f[dc]=pd.to_datetime(f[dc]);cp.append(pd.to_numeric(f.set_index(dc)[sid],errors='coerce').rename(sid))
    a=pd.concat(cp,axis=1).dropna().loc[:PLAN['market_snapshot_cutoff']];sig['cp_bill_spread']=(a.DCPF3M-a.DTB3)*100
    vol={k:ofr_series(sid) for k,sid in [('gcf_volume','REPO-GCF_TV_TOT-F'),('triparty_volume','REPO-TRIV1_TV_TOT-F')]}
    for k,s in vol.items():
        assert (s>0).all();save(s.rename('value').reset_index(),k+'.csv')
    # Five-observation contraction is a proxy, never a named investor's withdrawal.
    sig['triparty_withdrawal']=-np.log(vol['triparty_volume']).diff(5)
    sig={k:s.dropna().sort_index() for k,s in sig.items()}
    for k,s in sig.items():save(s.rename('value').reset_index().rename(columns={'observation_date':'date','index':'date'}),k+'_signal.csv')
    return sig,vol

def event_screen(s,mode,event_ids,key):
    if mode=='weekly':
        values=s.resample('W-FRI').max();end=pd.Timestamp(PLAN['market_snapshot_cutoff']);values=values[values.index<=end]
        # A source ending earlier in the week must not create a partial last week.
        values=values[values.index<=s.index.max()]
        threshold=values.shift().rolling(52,min_periods=26).quantile(.9)
    else:values=s;threshold=s.shift().rolling(252,min_periods=126).quantile(.9)
    f=pd.DataFrame({'value':values,'threshold':threshold});f['alarm']=(f.value>f.threshold).astype(float).where(f.threshold.notna()&f.value.notna())
    exclude=pd.Series(False,index=f.index);details=[]
    for eid in event_ids:
        a,b=map(pd.Timestamp,PLAN['events'][eid]);lo=a-pd.offsets.BDay(20);hi=b+pd.offsets.BDay(20)
        if mode=='weekly':
            exclude|=(f.index>=lo)&(f.index-pd.Timedelta(days=4)<=hi)
            records=[]
            for week,g in s.loc[a:b].groupby(s.loc[a:b].index.to_period('W-FRI')):
                t=week.end_time.normalize()
                if t in threshold.index and pd.notna(threshold.loc[t]):records.append((g.max(),threshold.loc[t]))
            alarms=[value>cutoff for value,cutoff in records];n=len(alarms);hit=bool(any(alarms)) if n else None
        else:
            exclude|=(f.index>=lo)&(f.index<=hi)
            alarms=f.loc[a:b].alarm.dropna().to_list();n=len(alarms);hit=bool(any(alarms)) if n>=2 else None
        details.append(dict(chain_id=key,mode=mode,event=eid,n_periods=n,hit=hit,alarm_fraction=float(np.mean(alarms)) if n else np.nan))
    valid=[x for x in details if x['hit'] is not None]
    hitrate=np.mean([x['hit'] for x in valid]) if valid else np.nan;quiet=f.loc[~exclude].alarm.mean();median=np.median([x['alarm_fraction'] for x in valid]) if valid else np.nan
    passed=len(valid)>=2 and hitrate>=2/3 and quiet<=.15 and (mode=='weekly' or median>=.25)
    save(f.reset_index().rename(columns={'index':'date','observation_date':'date'}),key+'_event_alerts.csv')
    return dict(chain_id=key,mode=mode,events_available=len(valid),events_hit=sum(x['hit'] for x in valid),hit_fraction=hitrate,non_event_alert_proxy=quiet,median_event_alarm_fraction=median,status='PASS_COMPONENT' if passed else ('FAIL_COMPONENT' if len(valid)>=2 else 'INSUFFICIENT'),first=str(s.index.min().date()),last=str(s.index.max().date())),details

def quarter_peak(s):
    a=s.groupby(s.index.to_period('Q')).agg(['max','count']);return a['max'].where(a['count']>=45)
def mean_period(s,freq):
    a=s.groupby(s.index.to_period(freq)).agg(['mean','count']);return a['mean'].where(a['count']>=(45 if freq=='Q' else 15))
def consecutive_diff(s,log=False):
    f=np.log(s.where(s>0)) if log else s
    return f.diff().where(np.r_[False,np.diff(s.index.asi8)==1])
def block_ci(x,y,index,block):
    n=len(x);starts=[i for i in range(n-block+1) if np.all(np.diff(index.asi8[i:i+block])==1)]
    if n<max(12,block*2) or not starts:return np.nan,np.nan
    ids=(np.random.default_rng(SEED).choice(starts,(2000,int(np.ceil(n/block))))[:,:,None]+np.arange(block)).reshape(2000,-1)[:,:n]
    a=rankdata(x[ids],axis=1);b=rankdata(y[ids],axis=1);a-=a.mean(axis=1,keepdims=True);b-=b.mean(axis=1,keepdims=True);den=np.sqrt((a*a).sum(axis=1)*(b*b).sum(axis=1));v=np.divide((a*b).sum(axis=1),den,out=np.full(2000,np.nan),where=den>0)
    return np.nanquantile(v,[.025,.975])

def association(test_id,chain,x,y,freq='Q',lag=0,role='primary'):
    x=x.copy();x.index=x.index+lag
    z=pd.concat([x.rename('x'),y.rename('y')],axis=1).replace([np.inf,-np.inf],np.nan).dropna().sort_index()
    z=z.loc[:('2026Q2' if freq=='Q' else '2026-09')]
    r=rho(z.x.to_numpy(),z.y.to_numpy());lo,hi=block_ci(z.x.to_numpy(),z.y.to_numpy(),z.index,4 if freq=='Q' else 12)
    adjacent=pd.Series(np.r_[False,np.diff(z.index.asi8)==1],index=z.index)
    changes=z.diff().where(adjacent,axis=0).dropna()
    p=float(spearmanr(z.x,z.y).pvalue) if np.isfinite(r) else np.nan
    row=dict(test_id=test_id,chain_id=chain,role=role,frequency=freq,lag=lag,n=len(z),first=str(z.index.min()),last=str(z.index.max()),rho=r,ci_low=lo,ci_high=hi,ordinary_p=p,
       rho_changes=rho(changes.x.to_numpy(),changes.y.to_numpy()),changes_n=len(changes),
       rho_earlier=rho(z.loc[:('2024Q4' if freq=='Q' else '2024-12')].x.to_numpy(),z.loc[:('2024Q4' if freq=='Q' else '2024-12')].y.to_numpy()),
       rho_later=rho(z.loc[('2025Q1' if freq=='Q' else '2025-01'):].x.to_numpy(),z.loc[('2025Q1' if freq=='Q' else '2025-01'):].y.to_numpy()),
       component_threshold_met=bool(len(z)>=(12 if freq=='Q' else 36) and r>=.3 and lo>0),full_chain_validated=False)
    save(z.assign(period=z.index.astype(str)).reset_index(drop=True),test_id+f'_lag{lag}_pairs.csv')
    return row

def auction_data():
    raw=pd.DataFrame(json.loads((ROOT/'raw/treasury_auctions_all.json').read_text()))
    raw['auction_date']=pd.to_datetime(raw.competitiveAuctionCloseDate,errors='coerce');raw['issue_date']=pd.to_datetime(raw.issueDate,errors='coerce')
    raw['offered_amount_usd']=pd.to_numeric(raw.offeringAmount,errors='coerce')
    keep=raw.auction_date.between('2018-01-01',PLAN['new_source_request_cutoff'])&raw.issue_date.between('2018-01-01',PLAN['new_source_request_cutoff'])&raw.treasuryAuctionResultsPressReleaseFileName.ne('')&raw.offered_amount_usd.gt(0)
    a=raw.loc[keep].copy();keys=['cusip','auction_date','issue_date','securityType'];assert not a.duplicated(keys).any()
    fields=['cusip','auction_date','issue_date','securityType','securityTerm','offered_amount_usd','bidToCoverRatio','highYield','treasuryAuctionResultsPressReleaseFileName']
    save(a[fields].sort_values('issue_date'),'treasury_auction_observations.csv')
    month=a.groupby(a.issue_date.dt.to_period('M')).offered_amount_usd.sum().loc[:'2026-09']
    save(month.rename('offered_amount_usd').reset_index().rename(columns={'issue_date':'month'}),'auction_monthly.csv')
    dump(dict(response_records=len(raw),selected_records=len(a),server_date_predicate_ignored=True,analysis_date_field='competitiveAuctionCloseDate',month_alignment='issueDate',completed_results_required=True,security_types=a.securityType.value_counts().to_dict(),offering_amount_is_not_total_actual_award=True,last_full_month=str(month.index.max()),future_announcements_excluded=True),'results/auction_filter_audit.json')
    return month

def funding_hac(signal):
    s=signal.rename('signal');thr=s.shift().rolling(252,min_periods=126).quantile(.9)
    spread=(ofr_series('REPO-GCF_AR_OO-F')-ofr_series('REPO-DVP_AR_OO-F'))*100
    f=pd.concat([s,thr.rename('threshold'),spread.rename('outcome_bp')],axis=1).dropna()
    f['shock']=(f.signal>f.threshold).astype(int);f['month_end']=[int((d+pd.offsets.BDay()).month!=d.month) for d in f.index];f['quarter_end']=[int(q_end(d)) for d in f.index]
    model=sm.OLS(f.outcome_bp,sm.add_constant(f[['shock','month_end','quarter_end']])).fit(cov_type='HAC',cov_kwds={'maxlags':5})
    lo,hi=model.conf_int().loc['shock'];save(f.reset_index().rename(columns={'index':'date'}),'GCF_FUNDING_pairs.csv')
    return dict(test_id='GCF_FUNDING',chain_id='FICC-GSD-05',role='primary',frequency='D',lag=0,n=len(f),first=str(f.index.min().date()),last=str(f.index.max().date()),coefficient_bp=float(model.params['shock']),hac_p=float(model.pvalues['shock']),ci_low=float(lo),ci_high=float(hi),r_squared=float(model.rsquared),shock_days=int(f.shock.sum()),component_threshold_met=bool(model.params['shock']>0 and model.pvalues['shock']<.05),full_chain_validated=False)

def run():
    check_inputs();new=extract_pqd();mbsd_cleared=reconcile_mbsd_activity(new);funds=extract_fund_flows();sig,vol=build_signals();auction=auction_data()
    p=pd.read_csv(ROOT/'inputs/response_quarterly.csv');p['period']=pd.PeriodIndex(p.quarter,freq='Q')
    entity=lambda e:p[p.entity.eq(e)].set_index('period')
    field=lambda e,c:new[new.entity.eq(e)].assign(period=lambda d:pd.PeriodIndex(d.quarter,freq='Q')).set_index('period')[c].sort_index()
    reconciled=pd.read_csv(ROOT/'data/financial_guarantee_reconciliation.csv');reconciled['period']=pd.PeriodIndex(reconciled.quarter,freq='Q')
    mbsd_unsettled=reconciled[reconciled.entity.eq('MBSD')].set_index('period').chosen_usd
    events=[];details=[]
    for c in PLAN['event_candidates']:
        a,b=event_screen(sig[c['signal']],c['mode'],c['events'],c['chain']);events.append(a);details+=b
    save(pd.DataFrame(events),'event_test_summary.csv','results');save(pd.DataFrame(details),'event_test_details.csv','results')
    monthly=pd.read_csv(ROOT/'inputs/ficc_monthly.csv');monthly['period']=pd.PeriodIndex(monthly.reporting_month,freq='M')
    gs=monthly[monthly.division.eq('GSD')&monthly.metric_code.eq('settled')].set_index('period').value_raw_source_units.sort_index();assert not gs.index.duplicated().any()
    monthly_gcf=mean_period(vol['gcf_volume'],'M');monthly_tri=mean_period(vol['triparty_volume'],'M')
    specs=[
      ('MMF_GCF','FICC-GSD-07',-funds.mmf_repo_flow_million_usd,-consecutive_diff(mean_period(vol['gcf_volume'],'Q'),log=True),'Q','primary'),
      ('TRIPARTY_GCF','FICC-GSD-07',-consecutive_diff(monthly_tri,log=True),-consecutive_diff(monthly_gcf,log=True),'M','supplementary_proxy'),
      ('CROSS_MARGIN','FICC-GSD-09',quarter_peak(sig['ust10_abs_change']),entity('GSD').im_qoq,'Q','primary'),
      ('AUCTION_GSD','FICC-GSD-10',consecutive_diff(auction,log=True),consecutive_diff(gs,log=True),'M','primary'),
      ('TBA_PIPELINE','FICC-MBSD-06',mbsd_cleared,mbsd_unsettled,'Q','primary'),
      ('CONCENTRATION','ECS-NSCC-04',field('NSCC','top5_position_share'),field('NSCC','top5_im_share'),'Q','internal_diagnostic'),
      ('SLD_LIQUIDITY','ECS-NSCC-05',quarter_peak(sig['vix']),entity('NSCC').vm_peak_usd,'Q','primary'),
    ]
    results=[funding_hac(sig['funding'])]
    for tid,cid,x,y,freq,role in specs:
        results.append(association(tid,cid,x,y,freq,role=role))
        results.append(association(tid,cid,x,y,freq,lag=1,role='secondary_next_period'))
    # Data-quality sensitivity: exclude the two independently repaired23.1.2 cells.
    exclude=mbsd_cleared.copy();exclude.loc[pd.PeriodIndex(['2021Q3','2022Q3'],freq='Q')]=np.nan
    results.append(association('TBA_PIPELINE_EXCLUDE_REPAIRED','FICC-MBSD-06',exclude,mbsd_unsettled,role='data_quality_sensitivity'))
    out=pd.DataFrame(results);out['bh_q_primary']=np.nan
    primary=out.index[out.role.eq('primary')];pv=out.loc[primary].ordinary_p.where(out.loc[primary].ordinary_p.notna(),out.loc[primary].hac_p)
    assert pv.notna().all();out.loc[primary,'bh_q_primary']=multipletests(pv.to_numpy(),method='fdr_bh')[1]
    save(out,'association_results.csv','results')
    # Seasonality-aware throughput robustness, declared in the plan.
    a=pd.read_csv(ROOT/'data/AUCTION_GSD_lag0_pairs.csv');a['month']=pd.PeriodIndex(a.period,freq='M').month
    x=pd.concat([a[['x']],pd.get_dummies(a.month,prefix='month',drop_first=True,dtype=float)],axis=1)
    fit=sm.OLS(a.y,sm.add_constant(x)).fit(cov_type='HAC',cov_kwds={'maxlags':3})
    dump(clean(dict(n=len(a),coefficient=float(fit.params['x']),hac_p=float(fit.pvalues['x']),ci=fit.conf_int().loc['x'].tolist(),scope='aggregate throughput, not auction-specific disruption',seasonality='calendar-month dummy variables',full_chain_validated=False)),'results/auction_seasonality_robustness.json')
    coverage=[dict(series_id=k,n=len(s),first=str(s.index.min().date()),last=str(s.index.max().date()),source_type='frozen market snapshot or new OFR final-vintage data') for k,s in sig.items()]
    coverage += [dict(series_id=k,n=len(s),first=str(s.index.min().date()),last=str(s.index.max().date()),source_type='new OFR final-vintage transaction volume') for k,s in vol.items()]
    save(pd.DataFrame(coverage),'series_coverage.csv','results')
    dump(clean(dict(run_at_utc=datetime.now(timezone.utc).isoformat(),definition_version='v3',new_chains=31,event_screened_chains=len(events),association_specs=8,data_quality_sensitivities=1,association_chain_count=out.chain_id.nunique(),primary_association_tests=int(out.role.eq('primary').sum()),complete_chain_validation_count=0,scores_computed=False,old_files_modified=False,plan_sha256=hashlib.sha256((ROOT/'analysis_plan.json').read_bytes()).hexdigest(),freshness='see series_coverage.csv; market snapshot not relabeled current, OFR final data ends June2026',no_holdout=True)),'run_manifest.json')
    print(out[['test_id','role','lag','n','rho','coefficient_bp','ci_low','ci_high','bh_q_primary']].round(4).to_string(index=False))
    print(pd.DataFrame(events)[['chain_id','events_hit','events_available','status']].to_string(index=False))
    return out,pd.DataFrame(events)

if __name__=='__main__':run()
