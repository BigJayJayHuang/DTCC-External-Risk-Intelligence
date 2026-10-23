"""Build two editable reports and language-matched vector/raster workflow diagrams."""
from pathlib import Path
import csv,html,json,math,textwrap,hashlib
from PIL import Image,ImageDraw,ImageFont
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from content import CN,EN,SOURCES

ROOT=Path(__file__).resolve().parents[1]
PROJECT=ROOT.parents[1]
OUT=ROOT/'output';OUT.mkdir(parents=True,exist_ok=True)
CH_FONT='/System/Library/Fonts/Hiragino Sans GB.ttc'
EN_FONT='/System/Library/Fonts/Supplemental/Arial.ttf'
EN_BOLD='/System/Library/Fonts/Supplemental/Arial Bold.ttf'

class Diagram:
    def __init__(self,lang):
        self.lang=lang;self.w,self.h=1600,2110
        self.image=Image.new('RGB',(self.w,self.h),'white');self.d=ImageDraw.Draw(self.image)
        self.svg=[f'<svg xmlns="http://www.w3.org/2000/svg" width="1600" height="2110" viewBox="0 0 1600 2110"><rect width="1600" height="2110" fill="white"/>']
    def font(self,size,bold=False):return ImageFont.truetype(CH_FONT if self.lang=='CN' else (EN_BOLD if bold else EN_FONT),size)
    def text(self,x,y,text,size=32,bold=False,fill='#142236',anchor='mm'):
        f=self.font(size,bold)
        self.d.text((x,y),text,font=f,fill=fill,anchor=anchor)
        ta='middle' if anchor=='mm' else 'start'
        family='Hiragino Sans GB,Arial Unicode MS,sans-serif' if self.lang=='CN' else 'Arial,sans-serif'
        self.svg.append(f'<text x="{x}" y="{y}" dominant-baseline="middle" text-anchor="{ta}" font-family="{family}" font-size="{size}" font-weight="{700 if bold else 400}" fill="{fill}">{html.escape(text)}</text>')
    def wrap(self,text,width,size):
        lines=[]
        for paragraph in text.split('\n'):
            line=''
            units=list(paragraph) if self.lang=='CN' else paragraph.split(' ')
            for unit in units:
                candidate=line+(' ' if line and self.lang!='CN' else '')+unit
                if self.d.textlength(candidate,font=self.font(size))>width and line:
                    lines.append(line);line=unit
                else:line=candidate
            lines.append(line)
        return lines
    def box(self,x,y,w,h,title,body,pending=False):
        fill='#f5f7f9' if pending else '#edf4fc';stroke='#8190a0' if pending else '#295e93'
        self.d.rounded_rectangle((x,y,x+w,y+h),radius=16,fill=fill,outline=None if pending else stroke,width=3)
        if pending:
            for xx in range(x+16,x+w-16,18):
                for yy in [y,y+h]:self.d.line([(xx,yy),(min(xx+11,x+w-16),yy)],fill=stroke,width=3)
            for yy in range(y+16,y+h-16,18):
                for xx in [x,x+w]:self.d.line([(xx,yy),(xx,min(yy+11,y+h-16))],fill=stroke,width=3)
            for bounds,start,end in [((x,y,x+32,y+32),180,270),((x+w-32,y,x+w,y+32),270,360),((x,y+h-32,x+32,y+h),90,180),((x+w-32,y+h-32,x+w,y+h),0,90)]:self.d.arc(bounds,start,end,fill=stroke,width=3)
        dash=' stroke-dasharray="11 7"' if pending else ''
        self.svg.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="16" fill="{fill}" stroke="{stroke}" stroke-width="3"{dash}/>')
        self.text(x+w/2,y+36,title,34,True)
        lines=self.wrap(body,w-46,29)
        if 81+(len(lines)-1)*35+18>h-6:raise ValueError('Diagram box overflow: '+title+' '+str(lines))
        for i,line in enumerate(lines):self.text(x+w/2,y+81+i*35,line,29)
    def arrow(self,points,dashed=False):
        self.d.line(points,fill='#506780',width=4,joint='curve')
        a,b=points[-2],points[-1];angle=math.atan2(b[1]-a[1],b[0]-a[0]);length=16;spread=.48
        tri=[b,(b[0]-length*math.cos(angle-spread),b[1]-length*math.sin(angle-spread)),(b[0]-length*math.cos(angle+spread),b[1]-length*math.sin(angle+spread))]
        self.d.polygon(tri,fill='#506780')
        dash=' stroke-dasharray="11 8"' if dashed else ''
        self.svg.append(f'<polyline points="'+ ' '.join(f'{x},{y}' for x,y in points)+f'" fill="none" stroke="#506780" stroke-width="4"{dash}/><polygon points="'+' '.join(f'{x},{y}' for x,y in tri)+'" fill="#506780"/>')
    def save(self):
        self.svg.append('</svg>')
        self.image.save(OUT/f'Workflow_{self.lang}.png',dpi=(240,240))
        (OUT/f'Workflow_{self.lang}.svg').write_text('\n'.join(self.svg))

def draw(lang):
    g=Diagram(lang);cn=lang=='CN'
    def t(a,b):return a if cn else b
    g.text(800,45,t('从公开信息到业务影响评估','From public information to business assessment'),43,True)
    g.text(800,94,t('数据与历史检验已完成  业务评分仍需验证','Data and historical tests exist; business scoring needs validation'),28)
    g.box(60,135,460,155,t('公开新闻与通知','News and notices'),t('官方新闻 监管文件\n保留来源和发布时间','Official news and rules\nSources and publication times'))
    g.box(570,135,460,155,t('市场与外部信息','Markets and external data'),t('利率 利差 波动\n公开技术和行业资料','Rates, spreads, volatility\nPublic industry/technology data'))
    g.box(1080,135,460,155,t('业务与风险披露','Business disclosures'),t('PQD 财务与运营\n保留服务 期间和单位','PQD, financials, operations\nService, period and units'))
    for x in [290,800,1310]:g.arrow([(x,290),(x,315),(800,315),(800,335)])
    g.box(190,340,1220,115,t('共享数据库与来源版本','Shared database and source versions'),t('同源只采一次 文档与事件分层 保存原值 证据和时间','Fetch shared sources once; separate documents/events; preserve evidence and times'))
    g.arrow([(800,455),(800,488)])
    g.box(190,490,1220,115,t('清洗 去重 时间与单位对齐','Clean deduplicate and align'),t('缺失不填零 日季频率分开 仅使用当时可用信息','Missing is not zero; preserve frequencies and information available at the time'))
    g.arrow([(800,605),(800,625),(410,625),(410,652)])
    g.arrow([(800,605),(800,625),(1190,625),(1190,652)])
    g.box(60,655,700,165,t('事件与通道特征','Events and channel features'),t('复核事件 ID 与风险通道\n各通道分别计算外部冲击 S','Review event IDs and channels\nMeasure shock S by channel'))
    g.box(840,655,700,165,t('业务画像与已公开状态','Business profiles and available state'),t('区分业务 实体和流程\n核实暴露 E 与业务专属脆弱度 V','Separate business, entity and process\nVerify exposure E and vulnerability V'))
    g.arrow([(410,820),(410,842),(800,842),(800,862)])
    g.arrow([(1190,820),(1190,842),(800,842),(800,862)])
    g.box(190,865,1220,145,t('先定义逻辑链和具体传导路径','Define logic chains and transmission pathways first'),t('事件 → 通道 → 会员或资产 → 业务步骤 → 可观察结果\n逐个箭头记录证据 缺口与适用范围','Event → channel → member/asset → process → observable outcome\nAttach evidence, gaps and scope to every arrow'))
    g.arrow([(800,1010),(800,1020),(410,1020),(410,1040)])
    g.arrow([(800,1010),(800,1020),(1190,1020),(1190,1040)])
    g.box(60,1045,700,250,t('历史训练与指标验证','Historical training and validation'),t('Ridge 与基准模型比较\n事件检验 同季及下一季 PQD 关联\n检查多重检验 样本和时期稳定性\n已执行检验 不等于已验证传导','Ridge versus baseline models\nEvent tests; same/next-quarter PQD tests\nCheck multiplicity and time stability\nTests run ≠ transmission established'))
    g.box(840,1045,700,250,t('拟整合的日常业务评估','Proposed daily business assessment'),t('先核实相关性与真实暴露\n按通道研究 S × E × V\n仅采用有支持的规则与参数\n数据不足时暂不评分','Verify relevance and actual exposure\nTest S × E × V by channel\nUse only supported rules/parameters\nDo not score missing evidence'),True)
    g.arrow([(760,1170),(839,1170)],True)
    g.arrow([(1190,1295),(1190,1322),(800,1322),(800,1350)])
    g.box(190,1355,1220,165,t('事件 × 业务 × 影响维度','Event × business × outcome dimension'),t('分别报告财务负担 流动性 操作质量和商业机会\n展示方向 幅度 单位 时间窗与证据等级','Separate financial burden, liquidity, service quality and opportunity\nReport direction, magnitude, units, horizon and evidence status'),True)
    g.arrow([(800,1520),(800,1560)])
    g.box(190,1565,1220,145,t('跟踪真实业务反应','Track observed business responses'),t('与本业务基线比较 保留其他解释\n季度峰值不能直接定位事件日影响','Compare with business-specific baselines; retain alternative explanations\nQuarterly peaks do not identify event-day impacts'),True)
    g.arrow([(800,1710),(800,1750)])
    g.box(190,1755,1220,145,t('用新增未见数据择期再验证','Revalidate with genuinely new observations'),t('冻结定义 后续分期验证 控制过拟合\n每日采集不会自动重训练历史模型','Freeze definitions; validate later periods; control overfitting\nDaily collection does not automatically retrain the models'),True)
    g.arrow([(800,1900),(800,1940),(25,1940),(25,1170),(60,1170)],True)
    g.text(800,1983,t('反馈到验证 而不是用已看过的测试集反复调参','Feed new evidence into validation, not repeated tuning on a seen test set'),26)
    g.text(800,2045,t('蓝色 已有工作   灰色虚线 待整合或验证   整理日期 2026-10-15','Blue: existing work   Dashed gray: proposed/unvalidated   Prepared 2026-10-15'),25)
    g.save()

def style_run(run,size=10.5,bold=False):
    run.font.name='Arial';run.font.size=Pt(size);run.bold=bold;run.font.color.rgb=RGBColor(0,0,0)
    fonts=run._element.get_or_add_rPr().rFonts
    fonts.set(qn('w:eastAsia'),'Hiragino Sans GB')
    for theme in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
        if qn('w:'+theme) in fonts.attrib:del fonts.attrib[qn('w:'+theme)]

def paragraph(doc,text,size=10.5,style=None,after=6):
    p=doc.add_paragraph(style=style);style_run(p.add_run(text),size)
    p.paragraph_format.space_after=Pt(after);p.paragraph_format.line_spacing=1.12
    return p

def heading(doc,text,level=1):
    p=doc.add_paragraph(style=f'Heading {level}');style_run(p.add_run(text),15 if level==1 else 11.5,True)
    p.paragraph_format.space_before=Pt(0 if level==1 else 8);p.paragraph_format.space_after=Pt(8 if level==1 else 5)
    return p

def table(doc,headers,rows,lang,section_index):
    n=len(headers);widths={3:[1.42,2.73,2.68],5:[1.65,.65,.65,.65,2.23]}[n]
    if section_index==13:widths=[.95,3.58,2.30]
    t=doc.add_table(rows=1,cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for i,h in enumerate(headers):t.rows[0].cells[i].text=h
    for row in rows:
        cells=t.add_row().cells
        for i,value in enumerate(row):cells[i].text=str(value)
    for ri,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if ri==0:pr.append(OxmlElement('w:tblHeader'))
        for ci,cell in enumerate(row.cells):
            cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E5EBF2' if ri==0 else ('F7F9FC' if ri%2==0 else 'FFFFFF'));cp.append(shade)
            borders=OxmlElement('w:tcBorders')
            for name in ['top','left','bottom','right']:
                edge=OxmlElement('w:'+name);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:color'),'D9D9D9');borders.append(edge)
            cp.append(borders)
            margins=OxmlElement('w:tcMar')
            for name,v in [('top',65),('bottom',65),('left',80),('right',80)]:
                edge=OxmlElement('w:'+name);edge.set(qn('w:w'),str(v));edge.set(qn('w:type'),'dxa');margins.append(edge)
            cp.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1.05
                if n==5 and ci>0:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:style_run(run,9.3 if section_index!=13 else 9.1,ri==0)
    paragraph(doc,'',size=1,after=4)

def hyperlink(p,label,url):
    r=p.part.relate_to(url,RT.HYPERLINK,is_external=True)
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),r)
    run=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'254E77');pr.append(color)
    sz=OxmlElement('w:sz');sz.set(qn('w:val'),'18');pr.append(sz);run.append(pr)
    text=OxmlElement('w:t');text.text=label;run.append(text);h.append(run);p._p.append(h)

def report(lang,pages):
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
    sec.top_margin=Inches(.62);sec.bottom_margin=Inches(.62);sec.left_margin=Inches(.72);sec.right_margin=Inches(.72)
    sec.footer_distance=Inches(.25)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
        s=doc.styles[name];s.font.name='Arial';s.font.color.rgb=RGBColor(0,0,0)
        s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Hiragino Sans GB')
    # The bundled Word starter may include accent borders on Title/Subtitle.
    # Strip paragraph borders from all styles rather than inheriting a blue rule.
    for borders in list(doc.styles.element.iter(qn('w:pBdr'))):borders.getparent().remove(borders)
    doc.styles['Normal'].font.size=Pt(10.5)
    doc.core_properties.title='DTCC 项目全流程与业务影响评估' if lang=='CN' else 'DTCC Project Workflow and Business Impact Assessment'
    doc.core_properties.author='Capstone project team';doc.core_properties.subject='Data processing model validation and transmission pathways'
    for i,item in enumerate(pages,1):
        if i>1:doc.add_page_break()
        if i==1:
            p=doc.add_paragraph(style='Title');style_run(p.add_run(doc.core_properties.title),22,True)
            paragraph(doc,'数据处理 模型验证 风险传导与跨业务比较' if lang=='CN' else 'Data processing model validation and business specific transmission',11,after=8)
            paragraph(doc,'2026 年 10 月 15 日' if lang=='CN' else 'October 15, 2026',9,after=10)
        heading(doc,f'{i:02} '+item['title'])
        for text in item.get('paras',[]):paragraph(doc,text)
        if item.get('figure'):
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(4)
            r=p.add_run();r.add_picture(str(OUT/f'Workflow_{lang}.png'),width=Inches(6.65))
            drawing=r._r.find('.//'+qn('wp:docPr'))
            if drawing is not None:drawing.set('descr','项目从公开数据到业务评估的整体流程' if lang=='CN' else 'End-to-end workflow from public data to business assessment')
        if item.get('table'):table(doc,*item['table'],lang,i)
        for title,paras in item.get('subsections',[]):
            heading(doc,title,2)
            for text in paras:paragraph(doc,text,9.1 if title in ['本地证据文件','Local evidence files'] else 10.5,after=5)
        for text in item.get('paras_after',[]):paragraph(doc,text)
        if item.get('table2'):table(doc,*item['table2'],lang,i)
        if item.get('sources'):
            heading(doc,'官方业务与方法资料' if lang=='CN' else 'Official business and methods references',2)
            for code,label,url in SOURCES:
                p=doc.add_paragraph();p.paragraph_format.space_after=Pt(3);hyperlink(p,f'[{code}] {label}',url)
        if item.get('note'):paragraph(doc,item['note'],9,after=3)
    p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    style_run(p.add_run('Capstone  |  '),8)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)
    file=OUT/f'DTCC_End_to_End_Workflow_Report_{lang}.docx';doc.save(file)
    return file

def evidence():
    r=PROJECT/'risk_business_ml'
    models=list(csv.DictReader((r/'results/model_comparison.csv').open(encoding='utf-8-sig')))
    augmented=[x for x in models if x['model']=='controls_plus_risks']
    assert len(augmented)==13 and sum(float(x['mae_improvement_vs_controls_pct'])>0 for x in augmented)==5
    expanded=json.loads((r/'risk_category_expansion/results/expanded_analysis_summary.json').read_text())
    assert expanded['same_and_lagged_associations_estimated']==482 and expanded['passing_exploratory_fdr_5pct']==0
    supplemental=json.loads((r/'feedback_revision_report_2026_10_06/results/analysis_summary.json').read_text())
    assert supplemental['comparisons']==882 and supplemental['counts']['all']['ordinary_fdr_5pct']==48 and supplemental['counts']['earlier']['ordinary_fdr_5pct']==0
    rows=list(csv.DictReader((r/'event_vulnerability_tests_2026_10_06/results/event_screen.csv').open(encoding='utf-8-sig')))
    assert len(rows)==7 and sum(x['event_gate_pass']=='True' for x in rows)==4
    events=[x for x in csv.DictReader((r/'event_vulnerability_tests_2026_10_06/results/event_detail.csv').open(encoding='utf-8-sig')) if x['event']=='svb_2023']
    peaks={x['indicator']:float(x['event_peak']) for x in events}
    assert peaks['ust10_abs_change']==23 and abs(peaks['baa_widen20']-41)<1e-8 and peaks['sofr_effr']==0 and peaks['vix']==26.52
    files=['results/model_comparison.csv','models/model_manifest.json','risk_category_expansion/results/expanded_analysis_summary.json',
           'event_vulnerability_tests_2026_10_06/analysis_plan.json','event_vulnerability_tests_2026_10_06/results/event_detail.csv',
           'event_vulnerability_tests_2026_10_06/results/event_screen.csv','feedback_revision_report_2026_10_06/results/analysis_summary.json']
    result={'prepared_date':'2026-10-15','frozen_model_observation_cutoff':'2026Q2','new_model_training':False,
            'ridge_targets':len(augmented),'ridge_targets_with_lower_holdout_mae':5,'event_candidates':7,'event_passes':4,
            'expanded_associations':482,'expanded_fdr_passes':0,'supplemental_records':882,'supplemental_full_sample_ordinary_bh_passes':48,
            'svb_window_market_peaks':peaks,'teaching_example':'Hypothetical S/E/V inputs; not fitted data or actual risk scores',
            'inputs':[{'project_relative_path':'risk_business_ml/'+str(p),'sha256':hashlib.sha256((r/p).read_bytes()).hexdigest()} for p in files]}
    (ROOT/'evidence_checks.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')

if __name__=='__main__':
    evidence()
    for lang,copy in [('CN',CN),('EN',EN)]:
        draw(lang);path=report(lang,copy);print(path)
