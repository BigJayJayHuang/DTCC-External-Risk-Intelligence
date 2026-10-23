"""Append all uncovered reference chains while preserving the complete v2 catalog."""
from pathlib import Path
from copy import deepcopy
from collections import Counter
from zipfile import ZipFile
import hashlib
import json
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn
from build_documents import ROOT, OUT, REFERENCE, digest, dump, run_style, table, xml_signature
from reference_additions_cn import TRANSLATIONS

QA = ROOT/'.qa/v3'
GROUPS = [('FICC','FICC'), ('ECS','ECS'), ('Asset Services','AssetSvc'), ('WIS','WIS')]

def put_paragraph(p, text):
    p.clear()
    run_style(p.add_run(text))

def before(anchor, text, style=None):
    p = anchor.insert_paragraph_before(style=style)
    run_style(p.add_run(text))
    return p

def add_chains(v2):
    existing = {c['chain_id'] for c in v2['chains']}
    missing = [(cid,row) for cid,row in v2['reference_catalog'].items() if cid not in existing]
    assert len(missing) == 31
    assert {cid for cid,row in missing} == set(TRANSLATIONS)
    result = []
    for cid, row in missing:
        cn = TRANSLATIONS[cid]
        indicator = row.get('Indicator to watch (optional)', '')
        result.append(dict(
            chain_id=cid, namespace='supplied_reference', legacy_ids=[],
            world_event={'CN':cn[0], 'EN':row['World event']},
            mechanism={'CN':cn[1], 'EN':row['Mechanism']},
            dtcc_process={'CN':cn[2], 'EN':row['DTCC process touched']},
            indicator_to_watch={'CN':cn[3], 'EN':indicator}, display_group=cid,
            introduced_in_version='v3', origin='previously_uncovered_supplied_reference',
            validation_status='not_tested_for_this_chain',
            data_readiness='not_assessed_for_this_chain',
            vulnerability_definition='not_selected',
            historical_results_inherited=False,
            source_reference_record=deepcopy(row)))
    return result

def create_doc(lang, v2, additions):
    source = OUT/f'Transmission_Chains_v2_{lang}.docx'
    doc = Document(source)
    section = deepcopy(doc.sections[0]._sectPr)
    old_tables = [xml_signature(t._tbl) for t in doc.tables[:2]]
    template = deepcopy(doc.tables[0]._tbl)
    title = '项目传导链完整目录' if lang == 'CN' else 'Capstone Transmission Chain Catalog'
    put_paragraph(doc.paragraphs[0], title)
    put_paragraph(doc.paragraphs[1], '2026 年 10 月 21 日  版本 v3' if lang == 'CN' else 'October 21, 2026  Version v3')
    opening_cn = [
        '本版覆盖 DTCC 提供文件中的全部 45 条示例链，并保留团队原有的 11 条独立扩展。原 v2 的 25 条不删减，补入此前未覆盖的 31 条；所有链继续采用外部事件、机制和具体流程的三步格式。',
        '完整目录共 56 条：FICC 23 条、ECS 26 条、资产服务 4 条、WIS 3 条。资产服务和 WIS 按参考文件的业务归属列示，不因为底层机构是 DTC 或 NSCC 就归入 ECS。原有 22 条假设的编号映射继续保留。',
    ]
    opening_en = [
        'This catalog covers all 45 examples in the supplied DTCC document and retains the 11 independent team extensions. All 25 v2 routes are preserved, with 31 previously uncovered reference chains added. Every chain retains the world-event, mechanism, and specific-process format.',
        'The 56 routes comprise 23 FICC, 26 ECS, 4 Asset Services, and 3 WIS entries. Asset Services and WIS retain their reference business classifications rather than being assigned to ECS because they use DTC or NSCC. The mappings for all 22 original hypotheses remain available.',
    ]
    for p, text in zip(doc.paragraphs[2:4], opening_cn if lang=='CN' else opening_en):
        put_paragraph(p, text)
    # The original 25 rows and original validation prose remain untouched.
    heading = '旧链到新链的完整对照' if lang=='CN' else 'Complete Legacy Mapping'
    anchor = next(p for p in doc.paragraphs if p.text == heading)
    before(anchor, '新增 DTCC 示例链' if lang=='CN' else 'Additional DTCC Reference Chains', 'Heading 1')
    before(anchor,
        '以下 31 条均为本轮新增，尚未按本链完成数据准备核对或检验。指标建议来自参考文件，不代表已经取得数据。相似事件、宏观信号或其他业务的既有检验，不自动转成这些链的通过结果。'
        if lang=='CN' else
        'All 31 entries below are newly added and have not undergone chain-specific data-readiness review or testing. Reference monitoring suggestions do not establish data availability. Similar events, macro signals, and tests of other businesses do not provide an inherited pass label.')
    for group,prefix in GROUPS:
        chains = [c for c in additions if c['chain_id'].startswith(prefix+'-')]
        label_cn = '资产服务' if prefix=='AssetSvc' else group
        before(anchor, f'{label_cn} 新增 {len(chains)} 条' if lang=='CN' else f'{group} Additions', 'Heading 2')
        rows = [[c['chain_id'],c['world_event'][lang],c['mechanism'][lang],c['dtcc_process'][lang]] for c in chains]
        headers=['编号','外部事件','机制','具体 DTCC 流程'] if lang=='CN' else ['ID','World event','Mechanism','DTCC process touched']
        new_table=table(doc,template,headers,rows,[1.625]*4)
        anchor._p.addprevious(new_table._tbl)
    # Only version references outside the two retained chain tables are updated.
    for p in doc.paragraphs:
        if lang=='CN' and p.text=='旧编号仅指 2026 年 10 月 16 日结果所用的 v1 编号。新编号对应本版，不可混用。':
            put_paragraph(p,'旧编号仅指 2026 年 10 月 16 日结果所用的 v1 编号。v2 的新编号在 v3 中继续适用；本轮新增链没有旧检验映射，不得按相同编号继承历史结果。')
        elif lang=='EN' and p.text.startswith('Legacy IDs refer only to the v1 definitions used in the October 16 results.'):
            put_paragraph(p,'Legacy IDs refer to the v1 definitions used in the October 16 results. Revised v2 IDs continue in v3. Newly added chains have no legacy-test mapping and must not inherit historical results through an ID match.')
        elif '配套 chain_registry_v2.json 保留' in p.text:
            put_paragraph(p,'配套 chain_registry_v3.json 保留完整目录与历史证据。新增链清单列出 31 条未检验链及参考监测建议。')
        elif p.text.startswith('The companion chain_registry_v2.json preserves'):
            put_paragraph(p,'chain_registry_v3.json contains the complete catalog and legacy evidence. The additions guides list all 31 untested chains and monitoring suggestions.')
    mapping_table=doc.tables[-1]
    put_paragraph(mapping_table.rows[0].cells[1].paragraphs[0], 'v3 编号' if lang=='CN' else 'Revised v3 ID')
    mapping_table.rows[0].cells[1].paragraphs[0].runs[0].bold=True
    assert [xml_signature(t._tbl) for t in doc.tables[:2]]==old_tables
    assert xml_signature(doc.sections[0]._sectPr)==xml_signature(section)
    assert [len(t.rows)-1 for t in doc.tables]==[10,15,13,11,4,3,22]
    doc.core_properties.title=title
    doc.core_properties.subject='Complete supplied reference coverage with retained team extensions'
    temp=QA/f'authored_v3_{lang}.docx'
    doc.save(temp)
    target=OUT/f'Transmission_Chains_v3_{lang}.docx'
    editable={'word/document.xml','docProps/core.xml'}
    with ZipFile(source) as base, ZipFile(temp) as authored, ZipFile(target,'w') as final:
        for info in base.infolist():
            final.writestr(info,authored.read(info.filename) if info.filename in editable else base.read(info.filename))
    with ZipFile(source) as base, ZipFile(target) as final:
        assert set(base.namelist())==set(final.namelist())
        assert all(base.read(n)==final.read(n) for n in base.namelist() if n not in editable)
    final_doc=Document(target)
    assert [xml_signature(t._tbl) for t in final_doc.tables[:2]]==old_tables
    return target

def write_guides(additions):
    for lang in ['CN','EN']:
        title='新增链与监测指标' if lang=='CN' else 'New Reference Chains and Monitoring Indicators'
        lines=[f'# {title}', '',
            '2026-10-21 v3：新增 31 条，均未检验；监测建议不代表数据已取得。' if lang=='CN' else
            'October 21, 2026 v3: 31 additions, all untested; monitoring suggestions do not establish data availability.', '']
        for group,prefix in GROUPS:
            lines.extend(['## '+('资产服务' if lang=='CN' and prefix=='AssetSvc' else group),''])
            for c in additions:
                if not c['chain_id'].startswith(prefix+'-'): continue
                indicator=c['indicator_to_watch'][lang] or ('参考文件未列出。' if lang=='CN' else 'Not specified in the reference.')
                lines.extend([f"### {c['chain_id']}",'',
                    c['world_event'][lang]+' '+c['mechanism'][lang], '',
                    ('流程：' if lang=='CN' else 'Process: ')+c['dtcc_process'][lang], '',
                    ('可选监测：' if lang=='CN' else 'Optional monitoring: ')+indicator, ''])
        (OUT/f'New_Reference_Chains_v3_{lang}.md').write_text('\n'.join(lines),encoding='utf-8')

def main():
    QA.mkdir(parents=True,exist_ok=True)
    v2_path=OUT/'chain_registry_v2.json'
    v2=json.loads(v2_path.read_text())
    protected=[v2_path,OUT/'Transmission_Chains_v2_CN.docx',OUT/'Transmission_Chains_v2_EN.docx',REFERENCE]
    old_checks=json.loads((ROOT/'.qa/delivery_checks.json').read_text())
    assert digest(protected[1])==old_checks['CN_output_sha256']
    assert digest(protected[2])==old_checks['EN_output_sha256']
    protected_hashes={str(p):digest(p) for p in protected}
    dump(QA/'v2_hashes_before.json',protected_hashes)
    additions=add_chains(v2)
    chains=deepcopy(v2['chains'])+additions
    ids=[c['chain_id'] for c in chains]
    assert len(ids)==len(set(ids))==56
    assert set(v2['reference_catalog'])<=set(ids)
    assert {c['chain_id'] for c in chains if c['namespace']=='supplied_reference'}==set(v2['reference_catalog'])
    assert len([c for c in chains if c['namespace']=='team'])==11
    assert chains[:25]==v2['chains']
    group_counts=Counter(c['chain_id'].split('-')[0] if not c['chain_id'].startswith('TEAM-')
                         else c['chain_id'].split('-')[1] for c in chains)
    assert dict(group_counts)==dict(FICC=23,ECS=26,AssetSvc=4,WIS=3)
    registry=deepcopy(v2)
    registry.update(version='v3',previous_version='v2',chains=chains,
        newly_added_chain_ids=[c['chain_id'] for c in additions],
        reference_coverage=dict(reference_rows=45,covered_reference_rows=45,missing_reference_ids=[]),
        business_counts=dict(group_counts),
        safe_result_join_keys=['definition_version','chain_id'],
        status_note='All v2 rows and all legacy evidence are unchanged. All 31 reference additions are untested; no historical results are inherited through matching IDs.')
    dump(OUT/'chain_registry_v3.json',registry)
    (QA/'artifact.md').write_text((ROOT/'.qa/artifact.md').read_text()+'''

## v3 authorized addition
The user requests adding every uncovered chain in the supplied reference. Start from each
verified v2 DOCX, retaining its two original chain tables and all 22 legacy mapping records.
Replace title, opening scope/counts, version labels, and registry filename only. Insert four
new tables before the legacy mapping: 13 FICC, 11 ECS, 4 Asset Services, 3 WIS. All new rows
are untested and retain the reference English verbatim with a faithful Chinese translation.
All 45 reference IDs must be covered once; preserve the 11 TEAM entries and all 25 v2 routes.
Only word/document.xml and docProps/core.xml may change from v2. v2 deliverables remain unchanged.
''',encoding='utf-8')
    outputs=[create_doc(lang,v2,additions) for lang in ['CN','EN']]
    write_guides(additions)
    assert all(digest(p)==h for p,h in protected_hashes.items())
    baseline=json.loads((ROOT/'.qa/source_hashes_before.json').read_text())
    assert all(digest(p)==h for p,h in baseline.items())
    checks=dict(total_chains=56,reference_chains_covered=45,team_extensions_retained=11,new_additions=31,
        new_chains_all_untested=True,business_counts=dict(group_counts),
        all_25_v2_rows_unchanged=True,all_22_legacy_mappings_unchanged=True,
        numerical_results_recomputed=False,v2_files_unchanged=True,legacy_evidence_unchanged=True,
        output_files=[str(p) for p in outputs],render_review='pending')
    dump(QA/'delivery_checks.json',checks)
    print(json.dumps(checks,ensure_ascii=False,indent=2))

if __name__=='__main__':
    main()
