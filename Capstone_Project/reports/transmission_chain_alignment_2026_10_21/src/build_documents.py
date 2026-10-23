"""Create reference-derived Word revisions without changing frozen evidence."""
from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile
import csv
import hashlib
import json
import re
import shutil
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from content import CHAINS, MAPPINGS, INTRO, GUIDANCE, VALIDATION, DATA_GUIDANCE

ROOT = Path(__file__).resolve().parents[1]
PROJECT = ROOT.parents[1]
TESTS = PROJECT / 'risk_business_ml/transmission_chain_tests_2026_10_16'
REFERENCE = Path('/Users/shanpeiyuan/Downloads/Transmission Chain Format and example/Transmission Chain Format and example.docx')
INPUT = ROOT / 'inputs'
OUT = ROOT / 'output'
QA = ROOT / '.qa'
EXPECTED_REFERENCE_HASH = '17f5fdc2cb0580802061fb9e1c93ad1b95b8d852701b4560c78cde50404b5af9'

def digest(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()

def dump(path, obj):
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding='utf-8')

def xml_signature(node):
    return (node.tag, tuple(sorted(node.attrib.items())), (node.text or '').strip(),
            tuple(xml_signature(child) for child in node))

def source_records(lang):
    with (TESTS / f'results/chain_status_{lang}.csv').open(encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))

def prepare():
    for p in [INPUT, OUT, QA]:
        p.mkdir(parents=True, exist_ok=True)
    assert digest(REFERENCE) == EXPECTED_REFERENCE_HASH
    sources = [REFERENCE, TESTS/'inputs/FICC_Transmission_Chain_Hypotheses_v1.docx',
        TESTS/'inputs/ECS Transmission Chain Hypotheses v1.docx',
        TESTS/'results/chain_status_CN.csv', TESTS/'results/chain_status_EN.csv']
    protected = sources + sorted((TESTS/'results').glob('*.csv')) + sorted((TESTS/'results').glob('*.json'))
    manifest = {str(p): digest(p) for p in set(protected)}
    dump(QA/'source_hashes_before.json', manifest)
    for p in sources:
        destination = INPUT/p.name
        if destination.exists():
            assert digest(destination) == digest(p)
        else:
            shutil.copy2(p, destination)
    ref = Document(REFERENCE)
    reference_rows = {}
    for t in ref.tables[1:]:
        headers = [c.text for c in t.rows[0].cells]
        for row in t.rows[1:]:
            values = [c.text for c in row.cells]
            record = dict(zip(headers, values))
            assert record['ID'] not in reference_rows
            reference_rows[record['ID']] = record
    dump(INPUT/'reference_chain_catalog.json', reference_rows)
    assert len(reference_rows) == 45
    with ZipFile(REFERENCE) as z:
        parts = [{'part': n, 'bytes': len(z.read(n)), 'sha256': hashlib.sha256(z.read(n)).hexdigest(),
                  'scope': 'editable body' if n == 'word/document.xml' else
                  'editable descriptive metadata' if n == 'docProps/core.xml' else 'preserve byte for byte'}
                 for n in z.namelist()]
        dump(QA/'reference_package_inventory.json', parts)
    artifact = f'''# Transmission chain revision contract

Reference: {REFERENCE}
SHA-256: {EXPECTED_REFERENCE_HASH}
Reference render: /private/tmp/capstone_chain_alignment_2026_10_21/reference_render
Reference pages: 15, all inspected. One section. No drawings or content controls.

## Content authority
The user's explicit merge rule controls the new text. Compare the complete 22 legacy entries
against all 45 reference rows by event, mechanism, and process rather than ID. Similar chains
follow the reference; conflicting causal statements defer to it; independent mechanisms stay.
The source's optional indicator is metadata, not a fourth causal step.

## Page and paragraph system
US Letter 8.5 by 11 inches, portrait, one-inch margins, one column. Header and footer distances
are 0.5 inches; headers and footers are blank. Normal inherits Calibri 11 pt, 8 pt after,
line spacing 259/240. Heading 1: bold 16 pt, 18 pt before, 6 pt after, keep next.
Heading 2: bold 13 pt, 14 pt before, 5 pt after, keep next. Title: 28 pt, 4 pt after.
All headings black. Chinese glyphs use Hiragino Sans GB via run-level East Asian fonts.

## Components and slots
word/document.xml/w:body children other than w:sectPr are editable content slots.
Replace the example body with title, dated introduction, FICC catalog, ECS catalog,
legacy mapping, evidence-transfer notes, and data guidance. The original section is preserve-only.
The four-column source chain table is the row pattern for the main catalog: ID, world event,
mechanism, process. Main widths remain the reference's equal 1.625-inch columns.
The three-column source table pattern is reused for the legacy mapping, with widths
1.25, 1.6, and 3.65 inches so mapping explanations remain legible.
All source text slots are replaced or intentionally removed; no reference example is presented
as the team's own executed test. No source-only expansion is automatically included as a test.

## Permitted content-driven adaptations
Use the existing Title style for the new title. Repeated header rows and non-splitting body rows
prevent the source's split-row pagination. Content height is automatic, never a fixed fit.
The third mapping-column width and Chinese fallback font are readability adaptations.
No new images, headers, footers, fields, relationships, or visual theme are introduced.

## Package and fidelity gates
Package inventory is in reference_package_inventory.json. Only word/document.xml and
docProps/core.xml may change. All other parts and relationships are byte-for-byte preserved
by overlaying the authored body onto the retained reference package. Section XML must match.
No source DOCX or legacy result is modified. Render both final DOCX files and inspect every page.
All 22 legacy IDs must map; revised IDs are unique; reference IDs must exist in the supplied file.
Preserve all old numerical evidence in the registry; changes of target or scope never inherit
a pass label. Route records are not scored or relabeled as new test results.
'''
    (QA/'artifact.md').write_text(artifact, encoding='utf-8')
    return reference_rows

def run_style(run, bold=None):
    run.font.name = 'Calibri'
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    fonts.set(qn('w:eastAsia'), 'Hiragino Sans GB')
    for attr in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
        fonts.attrib.pop(qn('w:'+attr), None)

def para(doc, text, style=None, keep=None):
    p = doc.add_paragraph(style=style)
    if keep is not None:
        p.paragraph_format.keep_with_next = keep
    run_style(p.add_run(text))
    return p

def cell_text(cell, text, header=False):
    for child in list(cell._tc):
        if child.tag != qn('w:tcPr'):
            cell._tc.remove(child)
    p = OxmlElement('w:p')
    cell._tc.append(p)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.08
    paragraph.paragraph_format.keep_with_next = False
    run = paragraph.add_run(str(text))
    run_style(run, header)
    run.font.size = Pt(11)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table(doc, template, headers, rows, widths):
    node = deepcopy(template)
    for tr in list(node.findall(qn('w:tr'))):
        node.remove(tr)
    body = doc._element.body
    body.insert(len(body)-1, node)
    from docx.table import Table
    t = Table(node, doc._body)
    t.autofit = False
    grid = t._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn('w:w'), str(round(width*1440)))
    for i, values in enumerate([headers] + rows):
        row = t.add_row()
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if i == 0:
            trpr.append(OxmlElement('w:tblHeader'))
        for j, (cell, value) in enumerate(zip(row.cells, values)):
            cell.width = Inches(widths[j])
            cell_text(cell, value, header=i == 0)
            tcpr = cell._tc.get_or_add_tcPr()
            # Match the reference's light-grey column-header treatment.
            for shade in list(tcpr.findall(qn('w:shd'))):
                tcpr.remove(shade)
            if i == 0:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'F2F2F2')
                tcpr.append(shade)
    return t

def build(lang):
    assert (QA/'artifact.md').exists()
    assert digest(REFERENCE) == EXPECTED_REFERENCE_HASH
    doc = Document(REFERENCE)
    chain_template = deepcopy(doc.tables[1]._tbl)
    mapping_template = deepcopy(doc.tables[0]._tbl)
    body = doc._element.body
    section_before = deepcopy(body.sectPr)
    for child in list(body):
        if child.tag != qn('w:sectPr'):
            body.remove(child)
    title = 'FICC 与 ECS 传导链修订版' if lang == 'CN' else 'FICC and ECS Transmission Chains'
    para(doc, title, 'Title')
    para(doc, '2026 年 10 月 21 日  版本 v2' if lang == 'CN' else 'October 21, 2026  Version v2')
    for text in INTRO[lang]:
        para(doc, text)
    para(doc, '修订规则' if lang == 'CN' else 'Revision Rules', 'Heading 2')
    for title2, text in GUIDANCE[lang]:
        p=doc.add_paragraph()
        run_style(p.add_run(title2+'  '), True)
        run_style(p.add_run(text))
    for group, prefix in [('FICC','FICC'), ('ECS','ECS')]:
        para(doc, (group+' 修订链') if lang == 'CN' else (group+' Revised Chains'), 'Heading 1')
        chains = [c for c in CHAINS if c['chain_id'].startswith(prefix) or c['chain_id'].startswith('TEAM-'+prefix)]
        rows = [[c['chain_id'], c['world_event'][lang], c['mechanism'][lang], c['dtcc_process'][lang]] for c in chains]
        headers = ['编号','外部事件','机制','具体 DTCC 流程'] if lang == 'CN' else ['ID','World event','Mechanism','DTCC process touched']
        table(doc, chain_template, headers, rows, [1.625]*4)
    para(doc, '旧链到新链的完整对照' if lang == 'CN' else 'Complete Legacy Mapping', 'Heading 1')
    para(doc, '旧编号仅指 2026 年 10 月 16 日结果所用的 v1 编号。新编号对应本版，不可混用。' if lang == 'CN' else
         'Legacy IDs refer only to the v1 definitions used in the October 16 results. Revised IDs refer to this version; do not interchange them.')
    rows = [[m['legacy_id'], '\n'.join(m['revised_ids']), m['action'][lang]] for m in MAPPINGS]
    headers = ['旧 v1 编号','v2 编号','处理及原因'] if lang == 'CN' else ['Legacy v1 ID','Revised v2 ID','Decision and reason']
    table(doc, mapping_template, headers, rows, [1.25,1.6,3.65])
    para(doc, '已有检验如何沿用' if lang == 'CN' else 'Use of Existing Validation', 'Heading 1')
    for title2, text in VALIDATION[lang]:
        para(doc, title2, 'Heading 2')
        para(doc, text)
    para(doc, '监测指标与数据衔接' if lang == 'CN' else 'Monitoring and Data Linkage', 'Heading 1')
    for title2, text in DATA_GUIDANCE[lang]:
        para(doc, title2, 'Heading 2')
        para(doc, text)
    para(doc, '依据与版本记录' if lang == 'CN' else 'Sources and Version Record', 'Heading 1')
    sources_cn = [
        'DTCC 提供的 Transmission Chain Format and example：三步格式、现有十条及新增候选表。',
        '团队 FICC Transmission Chain Hypotheses v1 与 ECS Transmission Chain Hypotheses v1：原 22 条假设及数据定义。',
        'chain_status_CN.csv 与 chain_status_EN.csv：2026 年 10 月 16 日已执行检验和未完成项。',
        '配套 chain_registry_v2.json 保留 25 条新链、22 条映射、每条可选监测指标、全部旧状态字段及参考目录。更改定义不改变历史结果；新规格的验证须另行冻结并执行。',
    ]
    sources_en = [
        'DTCC-supplied Transmission Chain Format and example: the three-step format, ten current examples, and additional candidate table.',
        'Team FICC Transmission Chain Hypotheses v1 and ECS Transmission Chain Hypotheses v1: the original 22 hypotheses and data definitions.',
        'chain_status_CN.csv and chain_status_EN.csv: executed tests and incomplete items as of October 16, 2026.',
        'The companion chain_registry_v2.json preserves 25 revised routes, 22 mappings, optional indicators, all legacy status fields, and the reference catalog. Definition changes do not change historical results; revised test specifications require a separate freeze and execution.',
    ]
    for text in sources_cn if lang == 'CN' else sources_en:
        para(doc, text)
    assert xml_signature(doc.sections[0]._sectPr) == xml_signature(section_before)
    doc.core_properties.title = title
    doc.core_properties.subject = 'Revision of team chains against supplied transmission-chain reference'
    doc.core_properties.author = 'Capstone project team'
    temp = QA/f'authored_{lang}.docx'
    doc.save(temp)
    target = OUT/f'Transmission_Chains_v2_{lang}.docx'
    # Preserve untouched package parts byte for byte, including styles and relationships.
    editable = {'word/document.xml', 'docProps/core.xml'}
    with ZipFile(REFERENCE) as base, ZipFile(temp) as authored, ZipFile(target,'w') as final:
        for info in base.infolist():
            final.writestr(info, authored.read(info.filename) if info.filename in editable else base.read(info.filename))
    with ZipFile(REFERENCE) as base, ZipFile(target) as final:
        assert set(base.namelist()) == set(final.namelist())
        assert all(base.read(n) == final.read(n) for n in base.namelist() if n not in editable)
    check = Document(target)
    assert len(check.sections) == 1
    assert len(check.tables) == 3
    assert [len(t.rows)-1 for t in check.tables] == [10,15,22]
    texts = '\n'.join(c.text for t in check.tables[:2] for row in t.rows[1:] for c in row.cells)
    assert all(c['chain_id'] in texts for c in CHAINS)
    return target

def main():
    refs = prepare()
    old_cn = source_records('CN')
    old_en = source_records('EN')
    ids = {r['编号'] for r in old_cn}
    assert len(CHAINS) == 25 and len(MAPPINGS) == 22 and len(ids) == 22
    assert len({c['chain_id'] for c in CHAINS}) == 25
    assert {m['legacy_id'] for m in MAPPINGS} == ids
    assert {x for c in CHAINS for x in c['legacy_ids']} == ids
    assert sum(c['namespace'] == 'supplied_reference' for c in CHAINS) == 14
    assert all(c['chain_id'] in refs for c in CHAINS if c['namespace']=='supplied_reference')
    for m in MAPPINGS:
        assert {c['chain_id'] for c in CHAINS if m['legacy_id'] in c['legacy_ids']} == set(m['revised_ids'])
    registry = dict(version='v2', definition_revision_date='2026-10-21', evidence_snapshot_date='2026-10-16',
        scores_computed=False, statistical_tests_rerun=False,
        formula='risk priority = external shock S * internal vulnerability V',
        chains=CHAINS, mappings=MAPPINGS, legacy_status_CN=old_cn, legacy_status_EN=old_en,
        reference_catalog=refs,
        status_note='Legacy evidence is retained without assigning any inherited pass label to a changed target or scope.')
    dump(OUT/'chain_registry_v2.json', registry)
    paths = [build(lang) for lang in ['CN','EN']]
    baseline = json.loads((QA/'source_hashes_before.json').read_text())
    assert all(digest(p) == sha for p,sha in baseline.items())
    checks = dict(revised_routes=25, reference_aligned=14, team_extensions=11, legacy_entries_mapped=22,
        split_legacy_entries=[m['legacy_id'] for m in MAPPINGS if len(m['revised_ids'])>1],
        merged_legacy_entries=['ECS-NSCC-04','ECS-DTC-01'],
        legacy_files_unchanged=True, original_reference_unchanged=digest(REFERENCE)==EXPECTED_REFERENCE_HASH,
        all_preserve_only_package_parts_identical=True, numerical_results_recomputed=False,
        output_files=[str(p) for p in paths], render_review='pending')
    dump(QA/'delivery_checks.json', checks)
    print(json.dumps(checks, ensure_ascii=False, indent=2))

if __name__ == '__main__':
    main()
