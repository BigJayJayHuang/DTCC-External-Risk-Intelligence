"""Resume public original-file collection. Does not calculate business metrics."""
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from hashlib import sha256
from zipfile import ZipFile
import json, re, subprocess, shutil, collections, argparse, io
from pypdf import PdfReader
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
FINANCIAL='https://www.dtcc.com/legal/financial-statements'
POLICY='https://www.dtcc.com/legal/policy-and-compliance'
MONTH_QUARTER={'march':1,'june':2,'september':3,'december':4}

def expected(record):
    title=record['title']
    years=re.findall(r'\b20\d{2}\b',title)
    q=re.search(r'\bQ([1-4])\b',title,re.I)
    return {'entity': title.split()[0] if record['category']=='financial_statements' else 'FICC and NSCC',
            'year':int(years[0]) if years else None,'quarter':int(q[1]) if q else None,
            'period':f'{years[0]} Q{q[1]}' if years and q else ('Annual '+ '/'.join(years) if years else 'unknown')}

def inspect(path):
    blob=path.read_bytes()
    data={'bytes':len(blob),'sha256':sha256(blob).hexdigest()}
    if blob.startswith(b'%PDF-'):
        reader=PdfReader(path)
        pages=[page.extract_text() or '' for page in reader.pages[:min(5,len(reader.pages))]]
        text='\n'.join(pages)
        assert len(text.strip())>100,'First pages are not readable'
        data.update(actual_format='pdf',pages=len(reader.pages),first_pages_text=text[:5000])
        cover=' '.join(pages[:2])
        context=text
    elif blob.startswith(b'PK'):
        with ZipFile(path) as z:
            assert z.testzip() is None,'ZIP integrity failure'
            assert 'xl/workbook.xml' in z.namelist(),'Not an XLSX workbook'
        workbook=load_workbook(io.BytesIO(blob),read_only=True,data_only=True)
        cover_parts=[]
        data.update(actual_format='xlsx',sheet_names=workbook.sheetnames)
        for sheet in workbook.worksheets[:3]:
            for row in sheet.iter_rows(min_row=1,max_row=min(40,sheet.max_row or 40),max_col=min(12,sheet.max_column or 12),values_only=True):
                cover_parts.append(' | '.join(str(x) for x in row if x is not None))
        workbook.close()
        cover='\n'.join(cover_parts)
        context=cover
        data['cover_text']=cover[:5000]
    else:
        raise ValueError('Response is neither a PDF nor an XLSX file')
    flat=re.sub(r'[^A-Z]','',cover.upper())
    if 'PUBLICQUANTITATIVE' in flat or 'QUANTITATIVEDISCLOSURES' in flat or data['actual_format']=='xlsx':
        data['document_kind']='public_quantitative_disclosure'
        data['actual_entity']='FICC and NSCC'
    else:
        data['document_kind']='financial_statements'
        data['actual_entity']=next((code for words,code in [
            ('THEDEPOSITORYTRUSTCLEARINGCORPORATION','DTCC'),
            ('FIXEDINCOMECLEARINGCORPORATION','FICC'),
            ('NATIONALSECURITIESCLEARINGCORPORATION','NSCC'),
            ('DEPOSITORYTRUSTCOMPANY','DTC')] if words in flat),'unconfirmed')
    q=re.search(r'\bQ([1-4])\s*[-/]?\s*(20\d{2})\b',cover,re.I)
    if not q:q=re.search(r'\bQ([1-4])\s*[-/]?\s*(20\d{2})\b',context,re.I)
    if q:
        data.update(actual_year=int(q[2]),actual_quarter=int(q[1]),period_evidence=q[0])
    else:
        dates=list(re.finditer(r'\b(March|June|September|December)\s+(?:31|30),?\s+(20\d{2})',cover,re.I))
        date=max(dates,key=lambda x:(int(x[2]),MONTH_QUARTER[x[1].lower()])) if dates else None
        if date:data.update(actual_year=int(date[2]),actual_quarter=MONTH_QUARTER[date[1].lower()],period_evidence=date[0])
    data['cover_excerpt']=cover[:1500]
    return data

def collect(url,records,existing):
    exp=expected(records[0])
    old=existing.get(url)
    warnings=[]
    if old and (ROOT/old['file']).is_file() and sha256((ROOT/old['file']).read_bytes()).hexdigest()==old.get('sha256'):
        path=ROOT/old['file']
        stamp=old.get('collected_at_utc')
        reused=True
    else:
        path=ROOT/'collection_support'/'staging'/f'{sha256(url.encode()).hexdigest()[:16]}.part'
        path.parent.mkdir(parents=True,exist_ok=True)
        stamp=datetime.now(timezone.utc).isoformat()
        reused=False
        if not path.is_file() or not path.read_bytes().startswith((b'%PDF-',b'PK')):
            result=subprocess.run(['/usr/bin/curl','--fail','--location','--connect-timeout','15','--max-time','60',
                                   '--retry','1','--silent','--show-error',url,'--output',str(path)],capture_output=True,text=True)
            if result.returncode:return dict(download_url=url,listing_records=records,status='download_failed',error=result.stderr.strip(),collected_at_utc=stamp)
    try:
        data=inspect(path)
        for rec in records:
            e=expected(rec)
            if rec['format']!=data['actual_format']:warnings.append(f"Portal format {rec['format']} != actual {data['actual_format']} for {rec['title']}")
            if data['actual_entity']!='unconfirmed' and e['entity']!=data['actual_entity']:warnings.append(f"Portal entity {e['entity']} != actual {data['actual_entity']}")
            if data.get('actual_year') and e['year']!=data['actual_year']:warnings.append(f"Portal year {e['year']} != actual {data['actual_year']} for {rec['title']}")
            if data.get('actual_quarter') and e['quarter'] and e['quarter']!=data['actual_quarter']:warnings.append(f"Portal quarter {e['quarter']} != actual {data['actual_quarter']} for {rec['title']}")
        identity_verified=data['actual_entity']!='unconfirmed' and bool(data.get('actual_year'))
        if not identity_verified:warnings.append('Document identity or reporting year needs manual review')
        if reused:
            relative=old['file']
        else:
            quarter=data.get('actual_quarter',exp['quarter'])
            year=data.get('actual_year',exp['year'])
            if data['document_kind']=='public_quantitative_disclosure':
                category='quantitative_disclosures'
                stem=f'DTCC_PQD_{year}_Q{quarter}'
            else:
                category='financial_statements'
                period=(f'Q{quarter}_{year}' if exp['quarter'] else f'Annual_{year}_{year-1}')
                stem=f"{data['actual_entity']}_{period}"
            if warnings:category='flagged_source_records'
            relative=f"{category}/{stem}__{url.rsplit('/',1)[-1][:8]}.{data['actual_format']}"
            target=ROOT/relative
            target.parent.mkdir(exist_ok=True)
            shutil.move(path,target)
        return dict(download_url=url,listing_records=records,file=relative,collected_at_utc=stamp,reused_existing_file=reused,
                    status='downloaded_with_metadata_warnings' if warnings else 'downloaded_and_identity_verified',
                    metadata_warnings=warnings,identity_verified=identity_verified,**data)
    except Exception as error:
        return dict(download_url=url,listing_records=records,status='verification_failed',error=str(error),collected_at_utc=stamp)

def main():
    queue=json.loads((ROOT/'historical_download_queue.json').read_text())
    groups=collections.defaultdict(list)
    for record in queue:groups[record['download_url']].append(record)
    old_manifest=json.loads((ROOT/'source_manifest.json').read_text())
    existing={x['download_url']:x for x in old_manifest['files']}
    progress=ROOT/'collection_support'/'historical_collection_progress.json'
    if progress.is_file():
        for record in json.loads(progress.read_text()).get('files',[]):
            if record.get('file'):existing[record['download_url']]=record
    completed=[]
    with ThreadPoolExecutor(max_workers=5) as pool:
        future_map={pool.submit(collect,url,records,existing):url for url,records in groups.items()}
        for future in as_completed(future_map):
            record=future.result();completed.append(record)
            saved={'collection_date_local':'2026-10-05','timezone':'America/New_York',
                   'source_pages':[FINANCIAL,POLICY],'files':sorted(completed,key=lambda x:x['download_url'])}
            intermediate=progress.with_suffix('.tmp')
            intermediate.write_text(json.dumps(saved,ensure_ascii=False,indent=2))
            intermediate.replace(progress)
            print(json.dumps({'done':len(completed),'total':len(groups),'status':record['status'],
                              'file':record.get('file'),'warnings':record.get('metadata_warnings',[]),'error':record.get('error')},ensure_ascii=False),flush=True)
    print(json.dumps({'records':len(queue),'unique_urls':len(groups),'status_counts':dict(collections.Counter(x['status'] for x in completed))}),flush=True)

if __name__=='__main__':main()
